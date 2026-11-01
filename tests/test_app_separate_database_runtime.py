import json
import os
import sqlite3
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException
from fastapi.testclient import TestClient
from pydantic import ValidationError

import app
from database_runtime import DatabaseRuntime


class SeparateDatabaseApplicationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.env_names = ("COPILOT_CONTROL_DB", "COPILOT_TENANT_DB_ROOT", "COPILOT_CORPUS_DB",
                          "COPILOT_DB", "COPILOT_OBJECT_STORAGE_ROOT", "COPILOT_OBJECT_STORAGE_BACKEND",
                          "COPILOT_ENV", "ADMIN_TOKEN")
        self.old_env = {name: os.environ.get(name) for name in self.env_names}
        os.environ.pop("COPILOT_DB", None)
        os.environ["COPILOT_CONTROL_DB"] = str(self.root / "control.sqlite3")
        os.environ["COPILOT_TENANT_DB_ROOT"] = str(self.root / "tenants")
        os.environ["COPILOT_CORPUS_DB"] = str(self.root / "corpus.sqlite3")
        os.environ["COPILOT_OBJECT_STORAGE_ROOT"] = str(self.root / "objects")
        os.environ["COPILOT_OBJECT_STORAGE_BACKEND"] = "local"
        os.environ["COPILOT_ENV"] = "development"
        self.admin_token = "test-only-shared-corpus-import-token-2026"
        os.environ["ADMIN_TOKEN"] = self.admin_token
        app._database_runtime = None

        self.runtime = DatabaseRuntime.from_env()
        self.runtime.initialize()
        self.firm_a, _ = self.runtime.provision_firm("Firm A")
        self.firm_b, _ = self.runtime.provision_firm("Firm B")
        self.password = "LongUniqueTestPassphrase-2026!"
        self.user_id = self.runtime.create_initial_firm_owner(
            self.firm_a, "First Owner", "owner@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(self.user_id, self.firm_b, "lawyer")
        self.client = TestClient(app.app)

    def tearDown(self):
        self.client.close()
        try:
            self.runtime.close_current_thread()
        except Exception:
            pass
        for name, value in self.old_env.items():
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value
        self.temp.cleanup()

    def _mutate(self, method, url, **kwargs):
        csrf = self.client.cookies.get("lc_csrf")
        headers = dict(kwargs.pop("headers", {}) or {})
        headers["X-CSRF-Token"] = csrf or ""
        return getattr(self.client, method)(url, headers=headers, **kwargs)

    def test_concurrent_logins_for_one_identity_issue_independent_sessions(self):
        def login_once(_):
            with TestClient(app.app) as client:
                response = client.post(
                    "/login",
                    json={"email": "OWNER@example.test", "password": self.password},
                )
                if response.status_code != 200:
                    return response.status_code, None
                return response.status_code, response.cookies.get("lc_session")

        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(login_once, range(4)))

        self.assertEqual([status for status, _ in results], [200, 200, 200, 200])
        session_tokens = [token for _, token in results]
        self.assertTrue(all(session_tokens))
        self.assertEqual(len(set(session_tokens)), 4, "concurrent logins must not reuse session tokens")

        control = self.runtime.control_connection()
        stored = control.execute(
            "SELECT COUNT(*) FROM auth_sessions WHERE user_id=?",
            (self.user_id,),
        ).fetchone()[0]
        self.assertGreaterEqual(stored, 4)


    def test_firm_connection_requires_trusted_context_and_never_falls_back(self):
        token = app._active_firm_context.set(None)
        try:
            with self.assertRaises(HTTPException) as raised:
                app.conn()
            self.assertEqual(raised.exception.status_code, 401)
            self.assertTrue(app._separate_runtime_enabled())
        finally:
            app._active_firm_context.reset(token)

    def test_disabling_firm_revokes_cached_tenant_connection(self):
        from tenant_context import FirmContextError

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        token = app._active_firm_context.set(context)
        try:
            cached = app.conn()
            self.assertEqual(cached.execute("SELECT COUNT(*) FROM cases").fetchone()[0], 0)
            self.assertTrue(self.runtime.disable_firm(self.firm_a))
            with self.assertRaises(FirmContextError):
                app.conn()
            with self.assertRaises(sqlite3.ProgrammingError):
                cached.execute("SELECT 1")
        finally:
            app._active_firm_context.reset(token)

    def test_security_audit_exports_require_firm_admin(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        control = self.runtime.control_connection()
        control.execute(
            "UPDATE firm_memberships SET firm_role='lawyer' WHERE user_id=? AND firm_id=?",
            (self.user_id, self.firm_a),
        )
        control.commit()
        for method, path in (
            ("get", "/audit/security-events"),
            ("get", "/audit/security-events.csv"),
        ):
            response = getattr(self.client, method)(path)
            self.assertEqual(response.status_code, 403, f"{path}: {response.text}")

    def test_security_audit_csv_neutralizes_spreadsheet_formulas(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        control = self.runtime.firm_connection(
            self.runtime.resolve_context(self.user_id, self.firm_a)
        )
        control.execute(
            "INSERT INTO security_audit_events(firm_id,user_id,event) VALUES(?,?,?)",
            (self.firm_a, self.user_id, " =HYPERLINK(\"https://attacker.invalid\",\"open\")"),
        )
        control.commit()
        response = self.client.get("/audit/security-events.csv")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("' =HYPERLINK", response.text)
        self.assertNotIn("\r\n =HYPERLINK", response.text)

    def test_direct_user_creation_is_disabled_in_separate_runtime(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        response = self._mutate(
            "post", "/users",
            json={"name": "Direct User", "email": "direct-user@example.test",
                  "password": self.password, "role": "lawyer"},
        )
        self.assertEqual(response.status_code, 410, response.text)
        control = self.runtime.control_connection()
        self.assertIsNone(control.execute(
            "SELECT id FROM users WHERE email_normalized='direct-user@example.test'"
        ).fetchone())
        self.assertEqual(control.execute(
            "SELECT COUNT(*) FROM identity_invitations WHERE email_normalized='direct-user@example.test'"
        ).fetchone()[0], 0)

    def test_membership_suspension_audit_failure_keeps_membership_active(self):
        target_id = self.runtime.create_global_identity(
            "Suspension Audit Target", "suspension-audit-target@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(target_id, self.firm_a, "lawyer")
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        control = self.runtime.control_connection()
        control.execute("""
            CREATE TRIGGER reject_membership_suspension_audit BEFORE INSERT ON platform_security_events
            WHEN NEW.event_type='membership_suspended'
            BEGIN SELECT RAISE(ABORT, 'membership suspension audit unavailable'); END
        """)
        control.commit()
        response = self._mutate("post", f"/users/{target_id}/disable")
        self.assertEqual(response.status_code, 503, response.text)
        self.assertEqual(control.execute(
            "SELECT status FROM firm_memberships WHERE user_id=? AND firm_id=?",
            (target_id, self.firm_a),
        ).fetchone()["status"], "active")
        self.assertEqual(control.execute(
            "SELECT COUNT(*) FROM platform_security_events WHERE event_type='membership_suspended' AND target_user_id=?",
            (target_id,),
        ).fetchone()[0], 0)

    def test_direct_user_creation_cannot_bypass_invitation_even_if_audit_ledger_is_unavailable(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        control = self.runtime.control_connection()
        before_events = control.execute(
            "SELECT COUNT(*) FROM platform_security_events WHERE event_type='membership_created'"
        ).fetchone()[0]
        control.execute("""
            CREATE TRIGGER reject_membership_creation_audit BEFORE INSERT ON platform_security_events
            WHEN NEW.event_type='membership_created'
            BEGIN SELECT RAISE(ABORT, 'membership creation audit unavailable'); END
        """)
        control.commit()
        response = self._mutate(
            "post", "/users",
            json={"name": "Not Created", "email": "not-created@example.test",
                  "password": self.password, "role": "lawyer"},
        )
        self.assertEqual(response.status_code, 410, response.text)
        self.assertIsNone(control.execute(
            "SELECT id FROM users WHERE email_normalized='not-created@example.test'"
        ).fetchone())
        self.assertEqual(control.execute(
            "SELECT COUNT(*) FROM platform_security_events WHERE event_type='membership_created'"
        ).fetchone()[0], before_events)

    def test_direct_identity_and_membership_routes_fail_closed_without_writes(self):
        target_user = self.runtime.create_global_identity(
            "Second Firm User", "second-firm-user@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(target_user, self.firm_b, "admin")
        self.assertTrue(self.runtime.disable_firm(self.firm_a))
        admin_context = {"id": self.user_id, "firm_id": self.firm_a, "role": "admin"}

        with self.assertRaises(HTTPException) as create_error:
            app.create_user(
                app.UserCreateIn(
                    name="Must Not Exist",
                    email="must-not-exist@example.test",
                    password=self.password,
                    role="lawyer",
                ),
                admin_context,
            )
        self.assertEqual(create_error.exception.status_code, 410)

        with self.assertRaises(HTTPException) as membership_error:
            app.add_existing_user_membership(
                app.ExistingMembershipIn(email="second-firm-user@example.test", role="lawyer"),
                admin_context,
            )
        self.assertEqual(membership_error.exception.status_code, 410)

        control = self.runtime.control_connection()
        self.assertIsNone(
            control.execute(
                "SELECT id FROM users WHERE email_normalized=?",
                ("must-not-exist@example.test",),
            ).fetchone()
        )
        self.assertEqual(
            control.execute(
                "SELECT COUNT(*) FROM firm_memberships WHERE user_id=? AND firm_id=?",
                (target_user, self.firm_a),
            ).fetchone()[0],
            0,
        )

    def test_disabled_global_identity_cannot_receive_direct_membership_grant(self):
        firm_c, _ = self.runtime.provision_firm("Disabled Identity Firm")
        disabled_user_id = self.runtime.create_initial_firm_owner(
            firm_c, "Disabled Owner", "disabled-owner@example.test", app.hash_pw(self.password)
        )
        control = self.runtime.control_connection()
        control.execute(
            "UPDATE users SET status='disabled',updated_at=datetime('now') WHERE id=?",
            (disabled_user_id,),
        )
        control.commit()

        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        response = self._mutate(
            "post", "/memberships",
            json={"email": "disabled-owner@example.test", "role": "lawyer"},
        )
        self.assertEqual(response.status_code, 410, response.text)
        self.assertEqual(
            control.execute(
                "SELECT COUNT(*) FROM firm_memberships WHERE user_id=? AND firm_id=?",
                (disabled_user_id, self.firm_a),
            ).fetchone()[0],
            0,
        )

    def test_existing_session_loses_suspended_firm_access_immediately(self):
        target_id = self.runtime.create_global_identity(
            "Existing Session Target", "existing-session-target@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(target_id, self.firm_a, "lawyer")
        self.runtime.create_firm_membership(target_id, self.firm_b, "lawyer")
        target_client = TestClient(app.app)
        try:
            target_login = target_client.post(
                "/login",
                json={"email": "existing-session-target@example.test", "password": self.password},
            )
            self.assertEqual(target_login.status_code, 200, target_login.text)
            csrf = target_client.cookies.get("lc_csrf") or ""
            selected_a = target_client.post(
                "/session/firm", json={"firm_id": self.firm_a},
                headers={"X-CSRF-Token": csrf},
            )
            self.assertEqual(selected_a.status_code, 200, selected_a.text)
            self.assertEqual(target_client.get("/cases").status_code, 200)

            owner_login = self.client.post(
                "/login", json={"email": "owner@example.test", "password": self.password},
            )
            self.assertEqual(owner_login.status_code, 200, owner_login.text)
            suspended = self._mutate("post", f"/users/{target_id}/disable")
            self.assertEqual(suspended.status_code, 200, suspended.text)

            # Keep the same authenticated session and cached firm connection:
            # the next request must re-check membership instead of trusting cache.
            denied_a = target_client.get("/cases")
            self.assertEqual(denied_a.status_code, 403, denied_a.text)
            selected_b = target_client.post(
                "/session/firm", json={"firm_id": self.firm_b},
                headers={"X-CSRF-Token": csrf},
            )
            self.assertEqual(selected_b.status_code, 200, selected_b.text)
            self.assertEqual(target_client.get("/cases").status_code, 200)
        finally:
            target_client.close()

    def test_suspending_one_membership_preserves_other_firm_access(self):
        target_id = self.runtime.create_global_identity(
            "Multi-firm Target", "multi-firm-target@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(target_id, self.firm_a, "lawyer")
        self.runtime.create_firm_membership(target_id, self.firm_b, "lawyer")
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        disabled = self._mutate("post", f"/users/{target_id}/disable")
        self.assertEqual(disabled.status_code, 200, disabled.text)
        self.assertTrue(disabled.json()["membership_suspended"])
        control = self.runtime.control_connection()
        self.assertEqual(
            control.execute("SELECT status FROM users WHERE id=?", (target_id,)).fetchone()["status"],
            "active",
        )
        memberships = {
            row["firm_id"]: row["status"] for row in control.execute(
                "SELECT firm_id,status FROM firm_memberships WHERE user_id=?", (target_id,)
            ).fetchall()
        }
        self.assertEqual(memberships[self.firm_a], "suspended")
        self.assertEqual(memberships[self.firm_b], "active")
        suspended_event = control.execute(
            "SELECT actor_user_id,target_user_id,firm_id,event_type FROM platform_security_events "
            "WHERE event_type='membership_suspended' AND target_user_id=? ORDER BY id DESC LIMIT 1",
            (target_id,),
        ).fetchone()
        self.assertIsNotNone(suspended_event)
        self.assertEqual(suspended_event["actor_user_id"], self.user_id)
        self.assertEqual(suspended_event["firm_id"], self.firm_a)

        from fastapi.testclient import TestClient
        target_client = TestClient(app.app)
        try:
            target_login = target_client.post(
                "/login", json={"email": "multi-firm-target@example.test", "password": self.password}
            )
            self.assertEqual(target_login.status_code, 200, target_login.text)
            denied_a = target_client.post(
                "/session/firm", json={"firm_id": self.firm_a},
                headers={"X-CSRF-Token": target_client.cookies.get("lc_csrf") or ""},
            )
            self.assertEqual(denied_a.status_code, 403, denied_a.text)
            allowed_b = target_client.post(
                "/session/firm", json={"firm_id": self.firm_b},
                headers={"X-CSRF-Token": target_client.cookies.get("lc_csrf") or ""},
            )
            self.assertEqual(allowed_b.status_code, 200, allowed_b.text)
            self.assertEqual(target_client.get("/cases").status_code, 200)
        finally:
            target_client.close()

    def test_firm_admin_cannot_suspend_another_firms_membership(self):
        firm_b_only_user = self.runtime.create_global_identity(
            "Firm B Only User", "firm-b-admin-isolation@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(firm_b_only_user, self.firm_b, "lawyer")
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        denied = self._mutate("post", f"/users/{firm_b_only_user}/disable")
        self.assertEqual(denied.status_code, 404, denied.text)
        membership = self.runtime.control_connection().execute(
            "SELECT status FROM firm_memberships WHERE user_id=? AND firm_id=?",
            (firm_b_only_user, self.firm_b),
        ).fetchone()
        self.assertIsNotNone(membership)
        self.assertEqual(membership["status"], "active")
        self.assertEqual(
            self.runtime.control_connection().execute(
                "SELECT status FROM users WHERE id=?", (firm_b_only_user,)
            ).fetchone()["status"],
            "active",
        )

    def test_user_listing_is_scoped_to_active_firm_memberships(self):
        firm_b_only_user = self.runtime.create_global_identity(
            "Firm B Only User", "firm-b-only@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(firm_b_only_user, self.firm_b, "lawyer")
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        firm_a_users = self.client.get("/users")
        self.assertEqual(firm_a_users.status_code, 200, firm_a_users.text)
        self.assertFalse(any(row["id"] == firm_b_only_user for row in firm_a_users.json()))

        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        firm_b_users = self.client.get("/users")
        self.assertEqual(firm_b_users.status_code, 200, firm_b_users.text)
        self.assertTrue(any(row["id"] == firm_b_only_user for row in firm_b_users.json()))

    def test_direct_membership_creation_requires_invitation_acceptance(self):
        firm_c, _ = self.runtime.provision_firm("Firm C")
        existing_user_id = self.runtime.create_initial_firm_owner(
            firm_c, "Second Firm Owner", "second-owner@example.test", app.hash_pw(self.password)
        )
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        response = self._mutate(
            "post", "/memberships",
            json={"email": " SECOND-OWNER@EXAMPLE.TEST ", "role": "lawyer"},
        )
        self.assertEqual(response.status_code, 410, response.text)
        control = self.runtime.control_connection()
        self.assertIsNone(control.execute(
            "SELECT 1 FROM firm_memberships WHERE user_id=? AND firm_id=?",
            (existing_user_id, self.firm_a),
        ).fetchone())
        self.assertEqual(control.execute(
            "SELECT COUNT(*) FROM platform_security_events "
            "WHERE event_type='membership_created' AND target_user_id=? AND firm_id=?",
            (existing_user_id, self.firm_a),
        ).fetchone()[0], 0)
        memberships = control.execute(
            "SELECT firm_id,firm_role,status FROM firm_memberships WHERE user_id=? ORDER BY firm_id",
            (existing_user_id,),
        ).fetchall()
        self.assertEqual(
            [(row["firm_id"], row["firm_role"], row["status"]) for row in memberships],
            [(firm_c, "admin", "active")],
        )

    def test_successful_case_deletion_preserves_firm_scoped_audit_snapshot(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate(
            "post", "/cases",
            json={"title": "CASE_DELETION_AUDIT_SENTINEL", "client": "Synthetic client"},
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        junior_id = self.runtime.create_global_identity(
            "Deletion Test Junior", "deletion-junior@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(junior_id, self.firm_a, "junior")
        added = self._mutate(
            "post", f"/cases/{case_id}/members",
            json={"user_id": junior_id, "role": "junior"},
        )
        self.assertEqual(added.status_code, 200, added.text)

        task = self._mutate(
            "post", f"/cases/{case_id}/tasks", json={"title": "Deletion snapshot task"}
        )
        self.assertEqual(task.status_code, 200, task.text)
        deadline = self._mutate(
            "post", f"/cases/{case_id}/deadlines",
            json={"title": "Deletion snapshot deadline", "due": "2026-11-15", "note": ""},
        )
        self.assertEqual(deadline.status_code, 200, deadline.text)
        document = self._mutate(
            "post", f"/cases/{case_id}/docs",
            json={"name": "deletion-snapshot.txt", "text": "Synthetic document retained in audit counts"},
        )
        self.assertEqual(document.status_code, 200, document.text)

        junior_client = TestClient(app.app)
        try:
            junior_login = junior_client.post(
                "/login",
                json={"email": "deletion-junior@example.test", "password": self.password},
            )
            self.assertEqual(junior_login.status_code, 200, junior_login.text)
            denied = junior_client.delete(
                f"/cases/{case_id}",
                headers={"X-CSRF-Token": junior_client.cookies.get("lc_csrf") or ""},
            )
            self.assertEqual(denied.status_code, 404, denied.text)
        finally:
            junior_client.close()

        deleted = self._mutate("delete", f"/cases/{case_id}")
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertTrue(deleted.json()["deleted"])

        tenant = self.runtime.firm_connection(
            self.runtime.resolve_context(self.user_id, self.firm_a)
        )
        self.assertIsNone(tenant.execute(
            "SELECT id FROM cases WHERE id=?", (case_id,)
        ).fetchone())
        audit = self.client.get("/audit/case-deletions")
        self.assertEqual(audit.status_code, 200, audit.text)
        row = next(item for item in audit.json()["items"] if item["case_id"] == case_id)
        self.assertEqual(row["firm_id"], self.firm_a)
        self.assertEqual(row["actor_user_id"], self.user_id)
        self.assertEqual(row["actor_name"], "First Owner")
        self.assertEqual(row["case_snapshot"]["title"], "CASE_DELETION_AUDIT_SENTINEL")
        self.assertEqual(row["case_snapshot"]["firm_id"], self.firm_a)
        self.assertGreaterEqual(len(row["events_snapshot"]), 3)
        self.assertGreaterEqual(row["related_counts"]["tasks"], 1)
        self.assertGreaterEqual(row["related_counts"]["deadlines"], 1)
        self.assertGreaterEqual(row["related_counts"]["docs"], 1)


    def test_case_delete_audit_failure_rolls_back_entire_deletion(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate(
            "post", "/cases", json={"title": "Deletion rollback sentinel"}
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        tenant_db = self.runtime.firm_connection(context)
        tenant_db.execute("""
            CREATE TRIGGER reject_case_deletion_audit
            BEFORE INSERT ON case_deletion_audit
            BEGIN SELECT RAISE(ABORT, 'injected audit persistence failure'); END
        """)
        tenant_db.commit()

        response = self._mutate("delete", f"/cases/{case_id}")
        self.assertEqual(response.status_code, 503, response.text)
        self.assertIsNotNone(
            tenant_db.execute(
                "SELECT id FROM cases WHERE id=? AND firm_id=?", (case_id, self.firm_a)
            ).fetchone(),
            "a case must not be deleted if its required audit snapshot cannot persist",
        )
        self.assertEqual(
            tenant_db.execute(
                "SELECT COUNT(*) FROM case_deletion_audit WHERE case_id=? AND firm_id=?",
                (case_id, self.firm_a),
            ).fetchone()[0],
            0,
        )
        self.assertIsNotNone(
            tenant_db.execute(
                "SELECT id FROM events WHERE case_id=?", (case_id,)
            ).fetchone(),
            "related matter history must remain intact after failed deletion",
        )

    def test_case_detail_hydrates_identity_names_from_control_plane_view(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "Identity Hydration Matter"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        member_id = self.runtime.create_global_identity(
            "Matter Associate", "matter-associate@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(member_id, self.firm_a, "lawyer")
        added = self._mutate(
            "post", f"/cases/{case_id}/members", json={"user_id": member_id, "role": "associate"}
        )
        self.assertEqual(added.status_code, 200, added.text)
        task = self._mutate(
            "post", f"/cases/{case_id}/tasks", json={"title": "Assigned to associate", "assignee": member_id}
        )
        self.assertEqual(task.status_code, 200, task.text)
        note = self._mutate(
            "post", f"/cases/{case_id}/notes", json={"text": "Identity hydration note"}
        )
        self.assertEqual(note.status_code, 200, note.text)

        detail = self.client.get(f"/cases/{case_id}")
        self.assertEqual(detail.status_code, 200, detail.text)
        body = detail.json()
        self.assertTrue(any(
            row["user_id"] == member_id and row["name"] == "Matter Associate"
            for row in body["members"]
        ))
        self.assertTrue(any(
            row["assignee"] == member_id and row["assignee_name"] == "Matter Associate"
            for row in body["tasks"]
        ))
        self.assertTrue(any(
            row["author"] == "First Owner" and row["text"] == "Identity hydration note"
            for row in body["notes"]
        ))

    def test_case_detail_update_and_delete_are_isolated_between_firms(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate(
            "post", "/cases", json={"title": "Firm A protected matter"}
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        context_a = self.runtime.resolve_context(self.user_id, self.firm_a)
        db_a = self.runtime.firm_connection(context_a)
        self.assertEqual(
            db_a.execute("SELECT title FROM cases WHERE id=?", (case_id,)).fetchone()[0],
            "Firm A protected matter",
        )
        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)

        detail = self.client.get(f"/cases/{case_id}")
        self.assertEqual(detail.status_code, 404, detail.text)
        update = self._mutate(
            "patch", f"/cases/{case_id}", json={"stage": "closed"}
        )
        self.assertEqual(update.status_code, 404, update.text)
        deletion = self._mutate("delete", f"/cases/{case_id}")
        self.assertEqual(deletion.status_code, 404, deletion.text)

        self.assertEqual(
            db_a.execute("SELECT title FROM cases WHERE id=?", (case_id,)).fetchone()[0],
            "Firm A protected matter",
        )

    def test_child_writes_reject_a_case_id_from_another_firm(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate(
            "post", "/cases", json={"title": "Firm A child-write boundary"}
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        attempts = (
            ("post", f"/cases/{case_id}/notes", {"text": "cross-firm note"}),
            ("post", f"/cases/{case_id}/deadlines", {"title": "Cross-firm deadline", "due": "2027-01-15"}),
            ("post", f"/cases/{case_id}/fees", {"title": "Cross-firm fee", "amount": 100, "paid": 0}),
            ("post", f"/cases/{case_id}/sessions", {"date": "2027-01-15", "note": "cross-firm session"}),
        )
        for method, path, payload in attempts:
            with self.subTest(path=path):
                response = self._mutate(method, path, json=payload)
                self.assertEqual(response.status_code, 404, response.text)

        upload = self._mutate(
            "post", f"/cases/{case_id}/upload",
            files={"file": ("private.txt", b"private source text", "text/plain")},
        )
        self.assertEqual(upload.status_code, 404, upload.text)

        context_a = self.runtime.resolve_context(self.user_id, self.firm_a)
        context_b = self.runtime.resolve_context(self.user_id, self.firm_b)
        db_a = self.runtime.firm_connection(context_a)
        db_b = self.runtime.firm_connection(context_b)
        for table in ("notes", "deadlines", "fees", "sessions"):
            with self.subTest(table=table):
                self.assertEqual(db_a.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0], 0)
                self.assertEqual(db_b.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0], 0)
        self.assertEqual(
            db_a.execute("SELECT title FROM cases WHERE id=?", (case_id,)).fetchone()[0],
            "Firm A child-write boundary",
        )

    def test_upload_preserves_staged_object_when_finalization_fails(self):
        import io
        from docx import Document

        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "Finalize Failure Recovery"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("Confidential original contract text for failure recovery.")
        document.save(buffer)
        payload = buffer.getvalue()

        with patch.object(
            app.LocalPrivateObjectStore,
            "finalize",
            side_effect=app.ObjectStorageError("simulated destination verification failure"),
        ), patch.object(app, "_mfa_required_for_user", return_value=False):
            response = self._mutate(
                "post", f"/cases/{case_id}/upload",
                files={"file": ("recovery-test.docx", payload,
                                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )

        self.assertEqual(response.status_code, 503, response.text)
        staged_files = list((self.root / "objects").glob("firm-*/*.staged"))
        self.assertEqual(len(staged_files), 1, "failed finalization must preserve the staged recovery object")
        self.assertEqual(staged_files[0].read_bytes(), payload)

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        self.assertEqual(
            firm_db.execute("SELECT COUNT(*) FROM original_file_objects").fetchone()[0],
            0,
            "rolled-back metadata must not advertise the failed upload as ready",
        )

    def test_original_file_upload_download_integrity_and_cross_firm_denial(self):
        import hashlib
        import io
        from docx import Document

        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "Original File Test"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        buffer = io.BytesIO()
        source = Document()
        source.add_paragraph("Confidential original contract text for storage verification.")
        source.save(buffer)
        original = buffer.getvalue()
        with patch.object(app, "_mfa_required_for_user", return_value=False):
            uploaded = self._mutate(
                "post", f"/cases/{case_id}/upload",
                files={"file": ("original-contract.docx", original,
                                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        body = uploaded.json()
        self.assertTrue(body["original_file_retained"])
        document_id = body["document_id"]

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        # Keep a read snapshot open while another request writes its audit event.
        # WAL must allow the download/audit write to complete without blocking.
        if not firm_db.in_transaction:
            firm_db.execute("BEGIN")
        metadata = firm_db.execute(
            "SELECT object_id,sha256,byte_size,state FROM original_file_objects WHERE document_id=? AND firm_id=?",
            (document_id, self.firm_a),
        ).fetchone()
        self.assertIsNotNone(metadata)
        self.assertEqual(metadata["state"], "ready")
        self.assertEqual(metadata["sha256"], hashlib.sha256(original).hexdigest())
        self.assertEqual(metadata["byte_size"], len(original))
        self.assertTrue((self.root / "objects" / f"firm-{self.firm_a}" /
                         (metadata["object_id"] + ".blob")).is_file())

        downloaded = self.client.get(f"/cases/{case_id}/docs/{document_id}/original")
        self.assertEqual(downloaded.status_code, 200, downloaded.text)
        self.assertEqual(downloaded.content, original)
        self.assertEqual(downloaded.headers["content-type"], "application/octet-stream")
        self.assertEqual(downloaded.headers["cache-control"], "no-store")
        # End the older snapshot before checking the newly committed audit row.
        firm_db.commit()
        self.assertTrue(firm_db.execute(
            "SELECT 1 FROM security_audit_events WHERE firm_id=? AND user_id=? AND event=?",
            (self.firm_a, self.user_id, f"original_document_download:{document_id}"),
        ).fetchone())

        # A same-firm user without case membership must not read the original
        # binary or inspect/change its hold/deletion lifecycle.
        nonmember_id = self.runtime.create_global_identity(
            "Original File Nonmember", "original-file-nonmember@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(nonmember_id, self.firm_a, "lawyer")
        nonmember_client = TestClient(app.app)
        try:
            nonmember_login = nonmember_client.post(
                "/login",
                json={"email": "original-file-nonmember@example.test", "password": self.password},
            )
            self.assertEqual(nonmember_login.status_code, 200, nonmember_login.text)
            nonmember_csrf = nonmember_client.cookies.get("lc_csrf") or ""
            original_base = f"/cases/{case_id}/docs/{document_id}/original"
            denied_attempts = [
                ("get", original_base, None),
                ("post", original_base + "/legal-hold", {"reason": "unauthorized hold"}),
                ("post", original_base + "/legal-hold/release", {"reason": "unauthorized release"}),
                ("get", original_base + "/legal-hold/events", None),
                ("post", original_base + "/deletion-requests", {"reason": "unauthorized deletion"}),
                ("get", original_base + "/deletion-requests", None),
                ("post", original_base + "/deletion-requests/1/approve", {"reason": "unauthorized approval"}),
            ]
            for method, url, payload in denied_attempts:
                with self.subTest(method=method, url=url):
                    kwargs = {}
                    if payload is not None:
                        kwargs["json"] = payload
                    if method == "post":
                        kwargs["headers"] = {"X-CSRF-Token": nonmember_csrf}
                    response = getattr(nonmember_client, method)(url, **kwargs)
                    self.assertEqual(response.status_code, 404, response.text)
        finally:
            nonmember_client.close()
        firm_db.commit()
        unchanged = firm_db.execute(
            "SELECT legal_hold,state FROM original_file_objects WHERE document_id=? AND firm_id=?",
            (document_id, self.firm_a),
        ).fetchone()
        self.assertEqual(unchanged["legal_hold"], 0)
        self.assertEqual(unchanged["state"], "ready")
        self.assertEqual(
            firm_db.execute(
                "SELECT COUNT(*) FROM original_file_deletion_requests WHERE document_id=? AND firm_id=?",
                (document_id, self.firm_a),
            ).fetchone()[0],
            0,
        )

        # A case member who is not a firm admin must not be able to change legal holds.
        member_id = self.runtime.create_global_identity(
            "Original File Junior", "original-file-junior@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(member_id, self.firm_a, "lawyer")
        added_member = self._mutate(
            "post", f"/cases/{case_id}/members", json={"user_id": member_id, "role": "junior"}
        )
        self.assertEqual(added_member.status_code, 200, added_member.text)
        member_client = TestClient(app.app)
        try:
            member_login = member_client.post(
                "/login", json={"email": "original-file-junior@example.test", "password": self.password}
            )
            self.assertEqual(member_login.status_code, 200, member_login.text)
            member_csrf = member_client.cookies.get("lc_csrf") or ""
            denied_hold = member_client.post(
                f"/cases/{case_id}/docs/{document_id}/original/legal-hold",
                json={"reason": "unauthorized hold attempt"},
                headers={"X-CSRF-Token": member_csrf},
            )
            self.assertEqual(denied_hold.status_code, 403, denied_hold.text)
        finally:
            member_client.close()

        # A failed audit insert must abort the hold state change atomically.
        firm_db.execute("""
            CREATE TRIGGER reject_original_hold_audit
            BEFORE INSERT ON security_audit_events
            WHEN NEW.event LIKE 'original_document_legal_hold_set:%'
            BEGIN SELECT RAISE(ABORT, 'hold audit unavailable'); END
        """)
        firm_db.commit()
        failed_hold = self._mutate(
            "post",
            f"/cases/{case_id}/docs/{document_id}/original/legal-hold",
            json={"reason": "This hold must roll back"},
        )
        self.assertEqual(failed_hold.status_code, 503, failed_hold.text)
        firm_db.commit()
        self.assertEqual(
            firm_db.execute(
                "SELECT legal_hold FROM original_file_objects WHERE document_id=?", (document_id,)
            ).fetchone()["legal_hold"],
            0,
        )
        firm_db.execute("DROP TRIGGER reject_original_hold_audit")
        firm_db.commit()

        held = self._mutate(
            "post",
            f"/cases/{case_id}/docs/{document_id}/original/legal-hold",
            json={"reason": "Preserve original evidence for pending proceedings"},
        )
        self.assertEqual(held.status_code, 200, held.text)
        self.assertTrue(held.json()["legal_hold"])
        firm_db.commit()
        hold_metadata = firm_db.execute(
            "SELECT legal_hold,hold_reason,hold_set_by,hold_set_at FROM original_file_objects WHERE document_id=?",
            (document_id,),
        ).fetchone()
        self.assertEqual(hold_metadata["legal_hold"], 1)
        self.assertEqual(hold_metadata["hold_reason"], "Preserve original evidence for pending proceedings")
        self.assertEqual(hold_metadata["hold_set_by"], self.user_id)
        self.assertTrue(hold_metadata["hold_set_at"])

        with self.assertRaises(sqlite3.IntegrityError):
            firm_db.execute(
                "DELETE FROM original_file_objects WHERE document_id=?", (document_id,)
            )
        firm_db.rollback()

        released = self._mutate(
            "post",
            f"/cases/{case_id}/docs/{document_id}/original/legal-hold/release",
            json={"reason": "The preservation instruction has been withdrawn"},
        )
        self.assertEqual(released.status_code, 200, released.text)
        self.assertFalse(released.json()["legal_hold"])
        firm_db.commit()
        release_metadata = firm_db.execute(
            "SELECT legal_hold,hold_released_by,hold_released_at,hold_release_reason "
            "FROM original_file_objects WHERE document_id=?",
            (document_id,),
        ).fetchone()
        self.assertEqual(release_metadata["legal_hold"], 0)
        self.assertEqual(release_metadata["hold_released_by"], self.user_id)
        self.assertTrue(release_metadata["hold_released_at"])
        self.assertEqual(
            release_metadata["hold_release_reason"], "The preservation instruction has been withdrawn"
        )
        for event in (
            f"original_document_legal_hold_set:{document_id}",
            f"original_document_legal_hold_released:{document_id}",
        ):
            self.assertTrue(firm_db.execute(
                "SELECT 1 FROM security_audit_events WHERE firm_id=? AND user_id=? AND event=?",
                (self.firm_a, self.user_id, event),
            ).fetchone(), event)

        hold_events = firm_db.execute(
            "SELECT id,action,reason,actor_user_id,occurred_at FROM original_file_hold_events "
            "WHERE firm_id=? AND document_id=? ORDER BY id",
            (self.firm_a, document_id),
        ).fetchall()
        self.assertEqual([event["action"] for event in hold_events], ["set", "release"])
        self.assertEqual(
            [event["reason"] for event in hold_events],
            [
                "Preserve original evidence for pending proceedings",
                "The preservation instruction has been withdrawn",
            ],
        )
        self.assertEqual([event["actor_user_id"] for event in hold_events], [self.user_id, self.user_id])
        self.assertTrue(all(event["occurred_at"] for event in hold_events))
        history_response = self.client.get(
            f"/cases/{case_id}/docs/{document_id}/original/legal-hold/events"
        )
        self.assertEqual(history_response.status_code, 200, history_response.text)
        self.assertEqual(
            [event["reason"] for event in history_response.json()["events"]],
            [
                "Preserve original evidence for pending proceedings",
                "The preservation instruction has been withdrawn",
            ],
        )
        with self.assertRaises(sqlite3.IntegrityError):
            firm_db.execute(
                "UPDATE original_file_hold_events SET reason='rewritten history' WHERE id=?",
                (hold_events[0]["id"],),
            )
        firm_db.rollback()
        with self.assertRaises(sqlite3.IntegrityError):
            firm_db.execute(
                "DELETE FROM original_file_hold_events WHERE id=?",
                (hold_events[0]["id"],),
            )
        firm_db.rollback()

        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        denied = self.client.get(f"/cases/{case_id}/docs/{document_id}/original")
        self.assertEqual(denied.status_code, 404, denied.text)
        denied_history = self.client.get(
            f"/cases/{case_id}/docs/{document_id}/original/legal-hold/events"
        )
        self.assertEqual(denied_history.status_code, 404, denied_history.text)

    def test_original_file_routes_require_document_to_match_path_case(self):
        import hashlib
        from original_file_storage import LocalPrivateObjectStore

        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        case_one = self._mutate("post", "/cases", json={"title": "Original file owner case"})
        case_two = self._mutate("post", "/cases", json={"title": "Different original path case"})
        self.assertEqual(case_one.status_code, 200, case_one.text)
        self.assertEqual(case_two.status_code, 200, case_two.text)
        case_one_id, case_two_id = case_one.json()["id"], case_two.json()["id"]

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        original = b"synthetic original evidence bound to case one"
        _doc_id, document_id = app.save_doc(
            firm_db,
            {"id": self.user_id, "firm_id": self.firm_a, "name": "First Owner"},
            case_one_id,
            "case-one-evidence.txt",
            "Synthetic extracted text only",
            content_hash=hashlib.sha256(original).hexdigest(),
            content_type="text/plain",
            byte_size=len(original),
        )
        store = LocalPrivateObjectStore(self.root / "objects")
        staged = store.stage(self.firm_a, original)
        store.finalize(self.firm_a, staged.object_id)
        firm_db.execute(
            "INSERT INTO original_file_objects("
            "object_id,document_id,firm_id,storage_backend,sha256,byte_size,content_type,state"
            ") VALUES(?,?,?,'local',?,?,?,'ready')",
            (
                staged.object_id, document_id, self.firm_a, staged.sha256,
                staged.byte_size, "text/plain",
            ),
        )
        firm_db.commit()

        # This user can access both cases, but the document belongs only to case one.
        base = f"/cases/{case_two_id}/docs/{document_id}/original"
        for method, path, kwargs in (
            ("get", base, {}),
            ("post", base + "/legal-hold", {"json": {"reason": "mismatched path case"}}),
            ("post", base + "/legal-hold/release", {"json": {"reason": "mismatched path case"}}),
            ("get", base + "/legal-hold/events", {}),
            ("post", base + "/deletion-requests", {"json": {"reason": "mismatched path case"}}),
            ("get", base + "/deletion-requests", {}),
        ):
            response = self._mutate(method, path, **kwargs) if method == "post" else self.client.get(path)
            self.assertEqual(response.status_code, 404, f"{path}: {response.text}")

        self.assertEqual(
            firm_db.execute(
                "SELECT state FROM original_file_objects WHERE document_id=? AND firm_id=?",
                (document_id, self.firm_a),
            ).fetchone()["state"],
            "ready",
        )
        self.assertEqual(
            firm_db.execute(
                "SELECT legal_hold FROM original_file_objects WHERE document_id=? AND firm_id=?",
                (document_id, self.firm_a),
            ).fetchone()[0],
            0,
        )
        self.assertEqual(
            firm_db.execute(
                "SELECT COUNT(*) FROM original_file_deletion_requests WHERE document_id=?",
                (document_id,),
            ).fetchone()[0],
            0,
        )

    def test_original_file_hold_and_deletion_routes_cannot_cross_firm_boundaries(self):
        import hashlib
        from original_file_storage import LocalPrivateObjectStore

        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created_a = self._mutate("post", "/cases", json={"title": "Firm A Original Evidence"})
        self.assertEqual(created_a.status_code, 200, created_a.text)
        case_id = created_a.json()["id"]

        firm_a_context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_a_db = self.runtime.firm_connection(firm_a_context)
        original = b"firm A confidential original bytes"
        _, document_id = app.save_doc(
            firm_a_db,
            {"id": self.user_id, "firm_id": self.firm_a, "name": "First Owner"},
            case_id,
            "firm-a-evidence.txt",
            "Extracted text for test only",
            content_hash=hashlib.sha256(original).hexdigest(),
            content_type="text/plain",
            byte_size=len(original),
        )
        store = LocalPrivateObjectStore(self.root / "objects")
        staged = store.stage(self.firm_a, original)
        store.finalize(self.firm_a, staged.object_id)
        firm_a_db.execute(
            "INSERT INTO original_file_objects("
            "object_id,document_id,firm_id,storage_backend,sha256,byte_size,content_type,state"
            ") VALUES(?,?,?,'local',?,?,?,'ready')",
            (
                staged.object_id, document_id, self.firm_a, staged.sha256,
                staged.byte_size, "text/plain",
            ),
        )
        firm_a_db.commit()
        requested = self._mutate(
            "post",
            f"/cases/{case_id}/docs/{document_id}/original/deletion-requests",
            json={"reason": "Retain pending deletion request for cross-firm test"},
        )
        self.assertEqual(requested.status_code, 200, requested.text)
        request_id = requested.json()["request_id"]

        foreign_admin_id = self.runtime.create_global_identity(
            "Firm B Original Admin", "firm-b-original-admin@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(foreign_admin_id, self.firm_b, "admin")
        foreign_client = TestClient(app.app)
        try:
            foreign_login = foreign_client.post(
                "/login",
                json={"email": "firm-b-original-admin@example.test", "password": self.password},
            )
            self.assertEqual(foreign_login.status_code, 200, foreign_login.text)
            created_b = foreign_client.post(
                "/cases",
                json={"title": "Firm B Unrelated Matter"},
                headers={"X-CSRF-Token": foreign_client.cookies.get("lc_csrf") or ""},
            )
            self.assertEqual(created_b.status_code, 200, created_b.text)
            # Matching parent IDs prove that authorization depends on the selected
            # firm database and document linkage, not on a globally guessed integer.
            self.assertEqual(created_b.json()["id"], case_id)

            base = f"/cases/{case_id}/docs/{document_id}/original"
            csrf = foreign_client.cookies.get("lc_csrf") or ""
            attempts = [
                ("get", base, None),
                ("post", base + "/legal-hold", {"reason": "foreign hold attempt"}),
                ("post", base + "/legal-hold/release", {"reason": "foreign release attempt"}),
                ("get", base + "/legal-hold/events", None),
                ("post", base + "/deletion-requests", {"reason": "foreign deletion request"}),
                ("get", base + "/deletion-requests", None),
                (
                    "post",
                    base + f"/deletion-requests/{request_id}/approve",
                    {"reason": "foreign approval attempt"},
                ),
            ]
            for method, url, body in attempts:
                with self.subTest(method=method, url=url):
                    kwargs = {}
                    if body is not None:
                        kwargs["json"] = body
                    if method == "post":
                        kwargs["headers"] = {"X-CSRF-Token": csrf}
                    response = getattr(foreign_client, method)(url, **kwargs)
                    self.assertEqual(response.status_code, 404, response.text)
        finally:
            foreign_client.close()

        original_row = firm_a_db.execute(
            "SELECT legal_hold,state FROM original_file_objects WHERE document_id=? AND firm_id=?",
            (document_id, self.firm_a),
        ).fetchone()
        self.assertEqual(original_row["legal_hold"], 0)
        self.assertEqual(original_row["state"], "ready")
        self.assertEqual(
            firm_a_db.execute(
                "SELECT status FROM original_file_deletion_requests WHERE id=? AND firm_id=?",
                (request_id, self.firm_a),
            ).fetchone()["status"],
            "requested",
        )
        firm_b_context = self.runtime.resolve_context(foreign_admin_id, self.firm_b)
        firm_b_db = self.runtime.firm_connection(firm_b_context)
        self.assertEqual(
            firm_b_db.execute("SELECT COUNT(*) FROM original_file_objects").fetchone()[0],
            0,
        )
        self.assertEqual(
            firm_b_db.execute("SELECT COUNT(*) FROM original_file_hold_events").fetchone()[0],
            0,
        )
        self.assertEqual(
            firm_b_db.execute("SELECT COUNT(*) FROM original_file_deletion_requests").fetchone()[0],
            0,
        )

    def test_original_file_deletion_requires_distinct_admin_and_legal_hold_cancels_it(self):
        import hashlib
        from original_file_storage import LocalPrivateObjectStore

        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "Deletion Workflow Matter"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        original = b"original evidence for deletion workflow"
        _, document_id = app.save_doc(
            firm_db,
            {"id": self.user_id, "firm_id": self.firm_a, "name": "First Owner"},
            case_id,
            "deletion-workflow.txt",
            "Extracted evidence text",
            content_hash=hashlib.sha256(original).hexdigest(),
            content_type="text/plain",
            byte_size=len(original),
        )
        store = LocalPrivateObjectStore(self.root / "objects")
        staged = store.stage(self.firm_a, original)
        store.finalize(self.firm_a, staged.object_id)
        firm_db.execute(
            "INSERT INTO original_file_objects("
            "object_id,document_id,firm_id,storage_backend,sha256,byte_size,content_type,state"
            ") VALUES(?,?,?,'local',?,?,?,'ready')",
            (
                staged.object_id, document_id, self.firm_a, staged.sha256,
                staged.byte_size, "text/plain",
            ),
        )
        firm_db.commit()

        requested = self._mutate(
            "post",
            f"/cases/{case_id}/docs/{document_id}/original/deletion-requests",
            json={"reason": "Remove a duplicate original after export review"},
        )
        self.assertEqual(requested.status_code, 200, requested.text)
        request_id = requested.json()["request_id"]

        same_admin_approval = self._mutate(
            "post",
            f"/cases/{case_id}/docs/{document_id}/original/deletion-requests/{request_id}/approve",
            json={"reason": "Attempted self-approval"},
        )
        self.assertEqual(same_admin_approval.status_code, 403, same_admin_approval.text)

        second_id = self.runtime.create_global_identity(
            "Second Firm Admin", "second-deletion-admin@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(second_id, self.firm_a, "admin")
        second_client = TestClient(app.app)
        try:
            second_login = second_client.post(
                "/login",
                json={"email": "second-deletion-admin@example.test", "password": self.password},
            )
            self.assertEqual(second_login.status_code, 200, second_login.text)
            approved = second_client.post(
                f"/cases/{case_id}/docs/{document_id}/original/deletion-requests/{request_id}/approve",
                json={"reason": "Second admin confirms approved export and deletion"},
                headers={"X-CSRF-Token": second_client.cookies.get("lc_csrf") or ""},
            )
            self.assertEqual(approved.status_code, 200, approved.text)
        finally:
            second_client.close()

        # A new legal hold supersedes a pending approved purge and appends an audit event.
        held = self._mutate(
            "post",
            f"/cases/{case_id}/docs/{document_id}/original/legal-hold",
            json={"reason": "Preserve evidence for newly identified litigation"},
        )
        self.assertEqual(held.status_code, 200, held.text)
        listed = self.client.get(
            f"/cases/{case_id}/docs/{document_id}/original/deletion-requests"
        )
        self.assertEqual(listed.status_code, 200, listed.text)
        request_record = listed.json()["requests"][0]
        self.assertEqual(request_record["status"], "cancelled")
        self.assertEqual(
            [event["action"] for event in request_record["events"]],
            ["requested", "approved", "cancelled"],
        )
        self.assertEqual(
            request_record["events"][-1]["reason"], "Automatically cancelled by legal hold"
        )

        under_hold = self._mutate(
            "post",
            f"/cases/{case_id}/docs/{document_id}/original/deletion-requests",
            json={"reason": "Must be rejected while held"},
        )
        self.assertEqual(under_hold.status_code, 409, under_hold.text)
        self.assertTrue(store.verify(
            self.firm_a, staged.object_id,
            expected_sha256=staged.sha256, expected_size=staged.byte_size,
        ))

    def test_production_original_storage_requires_complete_acceptance_manifest(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "Production Evidence Gate"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        from docx import Document
        import io
        buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("Production evidence gate test content")
        document.save(buffer)

        original_env = {
            key: os.environ.get(key)
            for key in (
                "COPILOT_ENV",
                "COPILOT_OBJECT_STORAGE_BACKEND",
                "COPILOT_DATABASE_ACCEPTANCE_FILE",
            )
        }
        os.environ["COPILOT_ENV"] = "production"
        os.environ["COPILOT_OBJECT_STORAGE_BACKEND"] = "s3"
        os.environ.pop("COPILOT_DATABASE_ACCEPTANCE_FILE", None)
        try:
            with patch.object(app, "_mfa_required_for_user", return_value=False):
                response = self._mutate(
                    "post", f"/cases/{case_id}/upload",
                    files={"file": ("production-evidence-gate.docx", buffer.getvalue(),
                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                )
        finally:
            for key, value in original_env.items():
                if value is None:
                    os.environ.pop(key, None)
                else:
                    os.environ[key] = value

        self.assertEqual(response.status_code, 503, response.text)
        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        self.assertEqual(
            firm_db.execute("SELECT COUNT(*) FROM docs WHERE case_id=?", (case_id,)).fetchone()[0],
            0,
            "production uploads must not create metadata without accepted operational evidence",
        )

    def test_original_file_storage_fails_closed_in_production_until_backend_is_available(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "Production Storage Gate"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        from docx import Document
        import io
        buffer = io.BytesIO()
        document = Document()
        document.add_paragraph("Production gate test content")
        document.save(buffer)
        os.environ["COPILOT_ENV"] = "production"
        os.environ["COPILOT_OBJECT_STORAGE_BACKEND"] = "s3"
        try:
            with patch.object(app, "load_acceptance_manifest", return_value={"ready": False}), \
                 patch.object(app, "_mfa_required_for_user", return_value=False):
                response = self._mutate(
                    "post", f"/cases/{case_id}/upload",
                    files={"file": ("production-gate.docx", buffer.getvalue(),
                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                )
        finally:
            os.environ["COPILOT_ENV"] = "development"
            os.environ["COPILOT_OBJECT_STORAGE_BACKEND"] = "local"
        self.assertEqual(response.status_code, 503, response.text)
        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        self.assertEqual(firm_db.execute(
            "SELECT COUNT(*) FROM docs WHERE case_id=?", (case_id,)
        ).fetchone()[0], 0)

    def test_original_storage_access_rechecks_target_fingerprint_before_s3_client(self):
        config = {
            "COPILOT_ENV": "production",
            "COPILOT_OBJECT_STORAGE_BACKEND": "s3",
            "COPILOT_DATABASE_ACCEPTANCE_FILE": "/protected/accepted-database-evidence.json",
            "COPILOT_S3_BUCKET": "changed-bucket",
            "COPILOT_S3_REGION": "us-east-1",
            "COPILOT_S3_KMS_KEY_ID": "arn:aws:kms:us-east-1:123456789012:key/changed-key",
            "COPILOT_S3_ENDPOINT_URL": "",
            "COPILOT_S3_RECOVERY_BUCKET": "private-recovery-test",
            "COPILOT_S3_RECOVERY_REGION": "us-east-1",
            "COPILOT_S3_RECOVERY_KMS_KEY_ID": "arn:aws:kms:us-east-1:123456789012:key/recovery",
            "COPILOT_S3_RECOVERY_ENDPOINT_URL": "",
        }
        with patch.dict(os.environ, config, clear=False), \
             patch.object(app, "load_acceptance_manifest", return_value={"ready": False}) as load_manifest, \
             patch("boto3.client") as boto_client:
            with self.assertRaises(HTTPException) as raised:
                app._original_file_store()
        self.assertEqual(raised.exception.status_code, 503)
        load_manifest.assert_called_once()
        self.assertEqual(
            load_manifest.call_args.kwargs.get("expected_storage_config_sha256"),
            app.original_storage_config_fingerprint(config),
        )
        boto_client.assert_not_called()

    def test_production_original_storage_requires_complete_acceptance_and_s3_backend(self):
        sentinel = object()
        config = {
            "COPILOT_ENV": "production",
            "COPILOT_OBJECT_STORAGE_BACKEND": "s3",
            "COPILOT_DATABASE_ACCEPTANCE_FILE": "/protected/accepted-database-evidence.json",
            "COPILOT_S3_BUCKET": "private-originals-test",
            "COPILOT_S3_REGION": "us-east-1",
            "COPILOT_S3_KMS_KEY_ID": "arn:aws:kms:us-east-1:123456789012:key/test-key",
            "COPILOT_S3_ENDPOINT_URL": "",
            "COPILOT_S3_RECOVERY_BUCKET": "private-recovery-test",
            "COPILOT_S3_RECOVERY_REGION": "us-east-1",
            "COPILOT_S3_RECOVERY_KMS_KEY_ID": "arn:aws:kms:us-east-1:123456789012:key/recovery",
            "COPILOT_S3_RECOVERY_ENDPOINT_URL": "",
        }
        with patch.dict(os.environ, config, clear=False), \
             patch.object(app, "load_acceptance_manifest", return_value={"ready": True}), \
             patch("boto3.client") as boto_client, \
             patch.object(app, "S3PrivateObjectStore", return_value=sentinel) as store_factory:
            result = app._original_file_store()
        self.assertIs(result, sentinel)
        self.assertEqual([call.args[0] for call in boto_client.call_args_list], ["s3", "kms"])
        store_factory.assert_called_once()
        self.assertEqual(store_factory.call_args.kwargs["bucket"], "private-originals-test")
        self.assertEqual(store_factory.call_args.kwargs["kms_key_id"], config["COPILOT_S3_KMS_KEY_ID"])

        local_config = {
            "COPILOT_ENV": "production",
            "COPILOT_OBJECT_STORAGE_BACKEND": "local",
            "COPILOT_DATABASE_ACCEPTANCE_FILE": "/protected/accepted-database-evidence.json",
        }
        with patch.dict(os.environ, local_config, clear=False), \
             patch.object(app, "load_acceptance_manifest", return_value={"ready": True}):
            with self.assertRaises(HTTPException) as raised:
                app._original_file_store()
        self.assertEqual(raised.exception.status_code, 503)
        self.assertIn("requires the accepted S3/KMS backend", raised.exception.detail)

    def test_cross_firm_document_and_membership_writes_are_denied(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate(
            "post", "/cases", json={"title": "Firm A document boundary"}
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        context_a_before = self.runtime.resolve_context(self.user_id, self.firm_a)
        db_a_before = self.runtime.firm_connection(context_a_before)
        initial_member_count = db_a_before.execute(
            "SELECT COUNT(*) FROM members WHERE case_id=?", (case_id,)
        ).fetchone()[0]
        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)

        attempts = (
            ("post", f"/cases/{case_id}/docs", {"name": "secret.txt", "text": "firm A private evidence"}),
            ("post", f"/cases/{case_id}/members", {"user_id": self.user_id, "role": "lead"}),
            ("post", f"/cases/{case_id}/tasks", {"title": "Cross-firm task"}),
        )
        for method, path, payload in attempts:
            with self.subTest(path=path):
                response = self._mutate(method, path, json=payload)
                self.assertEqual(response.status_code, 404, response.text)

        context_a = self.runtime.resolve_context(self.user_id, self.firm_a)
        context_b = self.runtime.resolve_context(self.user_id, self.firm_b)
        db_a = self.runtime.firm_connection(context_a)
        db_b = self.runtime.firm_connection(context_b)
        self.assertEqual(db_a.execute("SELECT COUNT(*) FROM docs WHERE case_id=?", (case_id,)).fetchone()[0], 0)
        self.assertEqual(db_a.execute("SELECT COUNT(*) FROM tasks WHERE case_id=?", (case_id,)).fetchone()[0], 0)
        self.assertEqual(db_a.execute("SELECT COUNT(*) FROM members WHERE case_id=?", (case_id,)).fetchone()[0], initial_member_count)
        for table in ("docs", "tasks", "members"):
            with self.subTest(table=table):
                self.assertEqual(db_b.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0], 0)

    def test_email_allowlist_cannot_grant_platform_review_capability(self):
        previous = os.environ.get("LEGAL_REVIEWER_EMAILS")
        os.environ["LEGAL_REVIEWER_EMAILS"] = "reviewer@example.test"
        try:
            roles = self.runtime.control_connection().execute(
                "SELECT COUNT(*) FROM user_platform_roles WHERE user_id=?",
                (self.user_id,),
            ).fetchone()[0]
            self.assertEqual(roles, 0)
            self.assertFalse(app._user_can_review_legal({
                "id": self.user_id, "role": "admin", "email": "REVIEWER@example.test"
            }))
        finally:
            if previous is None:
                os.environ.pop("LEGAL_REVIEWER_EMAILS", None)
            else:
                os.environ["LEGAL_REVIEWER_EMAILS"] = previous

    def test_platform_reviewer_role_is_separate_and_revocation_is_immediate(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        corpus = self.runtime.legal_corpus_connection()
        corpus.execute(
            "INSERT INTO legal_sources(source_id,authority,source_page) VALUES(?,?,?)",
            ("platform-review-source", "Platform review test", "https://example.test/source"),
        )
        corpus.execute(
            "INSERT INTO legal_documents(document_id,source_id,official_title,document_type) VALUES(?,?,?,?)",
            ("platform-review-doc", "platform-review-source", "Review test document", "statute"),
        )
        corpus.execute(
            "INSERT INTO legal_document_versions(version_id,document_id,version_label,source_url) VALUES(?,?,?,?)",
            (777, "platform-review-doc", "v1", "https://example.test/source/v1"),
        )
        corpus.commit()
        payload = {"review_type": "source_authenticity", "outcome": "pending", "notes": "test event"}
        url = "/legal-review/versions/777/events"

        # Firm-owner/admin membership does not imply a platform reviewer role.
        denied = self._mutate("post", url, json=payload)
        self.assertEqual(denied.status_code, 403, denied.text)

        control = self.runtime.control_connection()
        control.execute(
            "INSERT INTO user_platform_roles(user_id,role_name) VALUES(?,?)",
            (self.user_id, "platform_reviewer"),
        )
        control.commit()
        allowed = self._mutate("post", url, json=payload)
        self.assertEqual(allowed.status_code, 201, allowed.text)

        control.execute(
            "DELETE FROM user_platform_roles WHERE user_id=? AND role_name='platform_reviewer'",
            (self.user_id,),
        )
        control.commit()
        denied_after_revocation = self._mutate("post", url, json=payload)
        self.assertEqual(denied_after_revocation.status_code, 403, denied_after_revocation.text)
        self.assertEqual(
            corpus.execute("SELECT COUNT(*) FROM legal_review_events WHERE version_id=777").fetchone()[0],
            1,
        )

    def test_dashboard_calendar_and_finance_are_isolated_between_firms(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate(
            "post", "/cases", json={"title": "FIRM_A_PRIVATE_DASHBOARD_SENTINEL"}
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        deadline = self._mutate(
            "post", f"/cases/{case_id}/deadlines",
            json={"title": "FIRM_A_PRIVATE_DEADLINE_SENTINEL", "due": "2026-11-10"},
        )
        self.assertEqual(deadline.status_code, 200, deadline.text)
        fee = self._mutate(
            "post", f"/cases/{case_id}/fees",
            json={"title": "FIRM_A_PRIVATE_FEE_SENTINEL", "amount": 1234, "paid": 0, "due": "2026-11-10"},
        )
        self.assertEqual(fee.status_code, 200, fee.text)

        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        dashboard = self.client.get("/dashboard")
        self.assertEqual(dashboard.status_code, 200, dashboard.text)
        self.assertNotIn("FIRM_A_PRIVATE", dashboard.text)
        calendar = self.client.get("/calendar.ics")
        self.assertEqual(calendar.status_code, 200, calendar.text)
        self.assertNotIn("FIRM_A_PRIVATE", calendar.text)
        finance = self.client.get("/finance/overview")
        self.assertEqual(finance.status_code, 200, finance.text)
        self.assertEqual(finance.json()["fee_count"], 0)
        self.assertEqual(finance.json()["total_amount"], 0)

    def test_conflict_check_does_not_disclose_other_firm_matter(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate(
            "post", "/cases",
            json={"title": "Firm A conflict sentinel", "opposing_party": "CONFIDENTIAL_FIRM_A_OPPOSING_PARTY"},
        )
        self.assertEqual(created.status_code, 200, created.text)

        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        response = self._mutate(
            "post", "/conflicts", json={"q": "CONFIDENTIAL_FIRM_A_OPPOSING_PARTY"}
        )
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertEqual(result["matches"], [])
        self.assertFalse(result["possible_conflict"])
        self.assertFalse(result["admin_review_required"])

    def test_conflict_check_returns_only_generic_warning_for_same_firm_hidden_case(self):
        owner_login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(owner_login.status_code, 200, owner_login.text)
        created = self._mutate(
            "post",
            "/cases",
            json={
                "title": "PRIVATE_HIDDEN_CONFLICT_CASE_TITLE",
                "opposing_party": "PRIVATE_HIDDEN_CONFLICT_OPPOSING_PARTY",
            },
        )
        self.assertEqual(created.status_code, 200, created.text)

        nonmember_id = self.runtime.create_global_identity(
            "Conflict Nonmember", "conflict-nonmember@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(nonmember_id, self.firm_a, "lawyer")
        nonmember_client = TestClient(app.app)
        try:
            login = nonmember_client.post(
                "/login",
                json={"email": "conflict-nonmember@example.test", "password": self.password},
            )
            self.assertEqual(login.status_code, 200, login.text)
            response = nonmember_client.post(
                "/conflicts",
                json={"q": "PRIVATE_HIDDEN_CONFLICT_OPPOSING_PARTY"},
                headers={"X-CSRF-Token": nonmember_client.cookies.get("lc_csrf") or ""},
            )
            self.assertEqual(response.status_code, 200, response.text)
            result = response.json()
            self.assertEqual(result["matches"], [])
            self.assertTrue(result["possible_conflict"])
            self.assertTrue(result["admin_review_required"])
            self.assertNotIn("PRIVATE_HIDDEN_CONFLICT_CASE_TITLE", response.text)
            # The response echoes the caller's own query once, but must not add
            # the matched opposing-party value to any result or warning field.
            self.assertEqual(
                response.text.count("PRIVATE_HIDDEN_CONFLICT_OPPOSING_PARTY"), 1
            )
            self.assertNotIn(
                "PRIVATE_HIDDEN_CONFLICT_OPPOSING_PARTY", result["warning"]
            )
        finally:
            nonmember_client.close()

    def test_platform_reviewer_role_guards_all_shared_corpus_review_reads(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        self.assertFalse(login.json()["can_review_legal"])
        self.assertFalse(self.client.get("/session").json()["can_review_legal"])
        private_case = self._mutate(
            "post", "/cases", json={"title": "PRIVATE_TENANT_CATALOG_SENTINEL"}
        )
        self.assertEqual(private_case.status_code, 200, private_case.text)
        private_case_id = private_case.json()["id"]
        private_doc = self._mutate(
            "post", f"/cases/{private_case_id}/docs",
            json={"name": "PRIVATE_TENANT_CATALOG_SENTINEL.txt",
                  "text": "PRIVATE_TENANT_CATALOG_SENTINEL confidential matter text"},
        )
        self.assertEqual(private_doc.status_code, 200, private_doc.text)
        corpus = self.runtime.legal_corpus_connection()
        corpus.execute(
            "INSERT INTO legal_sources(source_id,authority,source_page) VALUES(?,?,?)",
            ("review-read-source", "Review read test", "https://example.test/source"),
        )
        corpus.execute(
            "INSERT INTO legal_documents(document_id,source_id,official_title,document_type) VALUES(?,?,?,?)",
            ("review-read-doc", "review-read-source", "Review read document", "statute"),
        )
        corpus.execute(
            "INSERT INTO legal_document_versions(version_id,document_id,version_label,source_url) VALUES(?,?,?,?)",
            (778, "review-read-doc", "v1", "https://example.test/source/v1"),
        )
        corpus.commit()
        urls = [
            "/legal-review/queue",
            "/legal-review/coverage",
            "/legal-review/catalog?kind=sources",
            "/legal-review/versions/778/events",
        ]

        for url in urls:
            with self.subTest(url=url):
                denied = self.client.get(url)
                self.assertEqual(denied.status_code, 403, denied.text)

        control = self.runtime.control_connection()
        control.execute(
            "INSERT INTO user_platform_roles(user_id,role_name) VALUES(?,?)",
            (self.user_id, "platform_reviewer"),
        )
        control.commit()
        self.assertTrue(self.client.get("/session").json()["can_review_legal"])
        for url in urls:
            with self.subTest(url=url, role="platform_reviewer"):
                allowed = self.client.get(url)
                self.assertEqual(allowed.status_code, 200, allowed.text)
                self.assertNotIn(
                    "PRIVATE_TENANT_CATALOG_SENTINEL", allowed.text,
                    "shared-corpus review endpoints must never expose firm matter content",
                )

        control.execute(
            "DELETE FROM user_platform_roles WHERE user_id=? AND role_name='platform_reviewer'",
            (self.user_id,),
        )
        control.commit()
        self.assertFalse(self.client.get("/session").json()["can_review_legal"])
        for url in urls:
            with self.subTest(url=url, role="revoked"):
                denied = self.client.get(url)
                self.assertEqual(denied.status_code, 403, denied.text)

    def test_qualified_counsel_allowlist_ignores_malformed_and_oversized_ids(self):
        with patch.dict(os.environ, {
            "LEGAL_QUALIFIED_EGYPTIAN_COUNSEL_USER_IDS":
                "42, 42, 0, -7, ١٢, abc, 9223372036854775808, " + ("9" * 5000) + ", 43"
        }):
            self.assertEqual(app._qualified_egyptian_counsel_user_ids(), {42, 43})

    def test_platform_reviewer_role_guards_shared_corpus_review_writes(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        corpus = self.runtime.legal_corpus_connection()
        corpus.execute(
            "INSERT INTO legal_sources(source_id,authority,source_page) VALUES(?,?,?)",
            ("review-write-source", "Review write test", "https://example.test/source"),
        )
        corpus.execute(
            "INSERT INTO legal_documents(document_id,source_id,official_title,document_type) VALUES(?,?,?,?)",
            ("review-write-doc", "review-write-source", "Review write document", "statute"),
        )
        corpus.execute(
            "INSERT INTO legal_document_versions(version_id,document_id,version_label,source_url,"
            "legal_status,review_status,authenticity_verified,status_verified,amendment_history_verified) "
            "VALUES(?,?,?,?,?,?,?,?,?)",
            (779, "review-write-doc", "v1", "https://example.test/source/v1",
             "under_review", "pending_lawyer_review", 0, 0, 0),
        )
        corpus.commit()

        payload = {
            "review_type": "legal_status",
            "outcome": "pending",
            "notes": "Synthetic authorization regression test",
        }
        endpoint = "/legal-review/versions/779/events"
        denied = self._mutate("post", endpoint, json=payload)
        self.assertEqual(denied.status_code, 403, denied.text)

        control = self.runtime.control_connection()
        control.execute(
            "INSERT INTO user_platform_roles(user_id,role_name) VALUES(?,?)",
            (self.user_id, "platform_reviewer"),
        )
        control.commit()
        allowed = self._mutate("post", endpoint, json=payload)
        self.assertEqual(allowed.status_code, 201, allowed.text)
        self.assertFalse(allowed.json()["status_changed"])

        missing_evidence = self._mutate(
            "post", endpoint,
            json={"review_type": "legal_status", "outcome": "pass"},
        )
        self.assertEqual(missing_evidence.status_code, 422, missing_evidence.text)
        credential_evidence = self._mutate(
            "post", endpoint,
            json={
                "review_type": "legal_status",
                "outcome": "pass",
                "evidence_url": "https://reviewer:secret@example.test/status",
            },
        )
        self.assertEqual(credential_evidence.status_code, 422, credential_evidence.text)

        passing = self._mutate(
            "post", endpoint,
            json={
                "review_type": "source_authenticity",
                "outcome": "pass",
                "evidence_url": "https://example.test/gazette/issue-123",
                "notes": "Evidence URL recorded for authorized human review.",
            },
        )
        self.assertEqual(passing.status_code, 201, passing.text)
        self.assertFalse(passing.json()["status_changed"])

        lawyer_payload = {
            "review_type": "lawyer_approval",
            "outcome": "pass",
            "evidence_url": "https://example.test/gazette/lawyer-review",
            "notes": "Synthetic counsel qualification attestation test.",
        }
        os.environ.pop("LEGAL_QUALIFIED_EGYPTIAN_COUNSEL_USER_IDS", None)
        unqualified = self._mutate("post", endpoint, json=lawyer_payload)
        self.assertEqual(unqualified.status_code, 403, unqualified.text)
        with patch.dict(os.environ, {
            "LEGAL_QUALIFIED_EGYPTIAN_COUNSEL_USER_IDS": str(self.user_id)
        }):
            qualified = self._mutate("post", endpoint, json=lawyer_payload)
        self.assertEqual(qualified.status_code, 201, qualified.text)
        self.assertIn(
            f"[qualified_egyptian_counsel_attested:user_id={self.user_id}]",
            qualified.json()["event"]["notes"],
        )

        current = corpus.execute(
            "SELECT review_status,legal_status,authenticity_verified,status_verified,"
            "amendment_history_verified FROM legal_document_versions WHERE version_id=779"
        ).fetchone()
        self.assertEqual(
            tuple(current),
            ("pending_lawyer_review", "under_review", 0, 0, 0),
            "review-event insertion must never auto-promote verification or legal status",
        )
        history = self.client.get(endpoint)
        self.assertEqual(history.status_code, 200, history.text)
        self.assertEqual(len(history.json()["events"]), 3)
        self.assertEqual(history.json()["events"][0]["review_type"], "lawyer_approval")
        self.assertIn(
            "source_authenticity",
            [event["review_type"] for event in history.json()["events"]],
        )
        source_review = next(
            event for event in history.json()["events"]
            if event["review_type"] == "source_authenticity"
        )
        self.assertEqual(source_review["reviewer"], "First Owner")

        control.execute(
            "DELETE FROM user_platform_roles WHERE user_id=? AND role_name='platform_reviewer'",
            (self.user_id,),
        )
        control.commit()
        denied_after_revocation = self._mutate("post", endpoint, json=payload)
        self.assertEqual(denied_after_revocation.status_code, 403, denied_after_revocation.text)
        self.assertEqual(
            corpus.execute(
                "SELECT COUNT(*) FROM legal_review_events WHERE version_id=779"
            ).fetchone()[0],
            3,
            "revoked platform role must not append a review event",
        )

    def test_shared_registry_discovery_never_returns_private_firm_documents(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate(
            "post", "/cases", json={"title": "Private matter discovery isolation"}
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        private_doc = self._mutate(
            "post", f"/cases/{case_id}/docs",
            json={"name": "PRIVATE_FIRM_ONLY_DISCOVERY_SENTINEL.pdf",
                  "text": "Confidential private firm document contents for discovery isolation test."},
        )
        self.assertEqual(private_doc.status_code, 200, private_doc.text)

        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        response = self._mutate(
            "post", "/research/discover", json={"q": "PRIVATE_FIRM_ONLY_DISCOVERY_SENTINEL"}
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["scope"], "metadata_only")
        self.assertEqual(response.json()["results"], [])
        self.assertFalse(response.json()["found"])

        # Metadata discovery is not the only public search path: ordinary
        # retrieval with no case_id must also exclude other firms' matter chunks.
        for endpoint in ("/search", "/ask"):
            retrieved = self._mutate(
                "post", endpoint,
                json={"q": "PRIVATE_FIRM_ONLY_DISCOVERY_SENTINEL", "scope": "all_permitted"},
            )
            self.assertEqual(retrieved.status_code, 200, retrieved.text)
            self.assertNotIn("PRIVATE_FIRM_ONLY_DISCOVERY_SENTINEL", retrieved.text)

    def test_same_firm_nonmember_cannot_search_or_ask_private_case_documents(self):
        owner_login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(owner_login.status_code, 200, owner_login.text)
        created = self._mutate(
            "post", "/cases", json={"title": "Same-firm private search boundary matter"}
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        uploaded = self._mutate(
            "post", f"/cases/{case_id}/docs",
            json={
                "name": "same-firm-private-search.txt",
                "text": "SAME_FIRM_NONMEMBER_PRIVATE_SENTINEL " * 12,
            },
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)

        nonmember_id = self.runtime.create_global_identity(
            "Same Firm Nonmember",
            "same-firm-search-nonmember@example.test",
            app.hash_pw(self.password),
        )
        self.runtime.create_firm_membership(nonmember_id, self.firm_a, "lawyer")
        nonmember = TestClient(app.app)
        try:
            login = nonmember.post(
                "/login",
                json={"email": "same-firm-search-nonmember@example.test", "password": self.password},
            )
            self.assertEqual(login.status_code, 200, login.text)
            csrf = nonmember.cookies.get("lc_csrf") or ""
            headers = {"X-CSRF-Token": csrf}

            # Explicit matter scope must authorize the case before retrieval.
            for endpoint in ("/search", "/ask"):
                denied = nonmember.post(
                    endpoint,
                    json={
                        "q": "SAME_FIRM_NONMEMBER_PRIVATE_SENTINEL",
                        "case_id": case_id,
                        "scope": "firm_documents",
                    },
                    headers=headers,
                )
                self.assertEqual(denied.status_code, 404, f"{endpoint}: {denied.text}")
                self.assertNotIn("SAME_FIRM_NONMEMBER_PRIVATE_SENTINEL", denied.text)

            # Broad scope means all *permitted* matter context, not all documents
            # owned by the firm; a same-firm nonmember still cannot retrieve this case.
            for endpoint in ("/search", "/ask"):
                response = nonmember.post(
                    endpoint,
                    json={
                        "q": "SAME_FIRM_NONMEMBER_PRIVATE_SENTINEL",
                        "scope": "all_permitted",
                    },
                    headers=headers,
                )
                self.assertEqual(response.status_code, 200, f"{endpoint}: {response.text}")
                self.assertNotIn("SAME_FIRM_NONMEMBER_PRIVATE_SENTINEL", response.text)
        finally:
            nonmember.close()

    def test_search_and_ask_reject_case_ids_from_another_firm(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        created_case = self._mutate(
            "post", "/cases", json={"title": "Firm A confidential search matter"}
        )
        self.assertEqual(created_case.status_code, 200, created_case.text)
        case_id = created_case.json()["id"]
        uploaded = self._mutate(
            "post", f"/cases/{case_id}/docs",
            json={"name": "search-private.txt", "text": "FIRM_A_ONLY_SEARCH_SENTINEL " * 12},
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)

        context_a = self.runtime.resolve_context(self.user_id, self.firm_a)
        db_a = self.runtime.firm_connection(context_a)
        self.assertEqual(
            db_a.execute("SELECT COUNT(*) FROM docs WHERE case_id=?", (case_id,)).fetchone()[0],
            1,
        )
        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)

        for endpoint in ("/search", "/ask"):
            response = self._mutate(
                "post", endpoint,
                json={"q": "FIRM_A_ONLY_SEARCH_SENTINEL", "case_id": case_id,
                      "scope": "all_permitted"},
            )
            self.assertEqual(response.status_code, 404, f"{endpoint}: {response.text}")
            self.assertNotIn("FIRM_A_ONLY_SEARCH_SENTINEL", response.text)

        # Cross-firm attempts cannot change the source firm's case or document.
        self.assertEqual(
            db_a.execute("SELECT COUNT(*) FROM cases WHERE id=?", (case_id,)).fetchone()[0],
            1,
        )
        self.assertEqual(
            db_a.execute("SELECT COUNT(*) FROM docs WHERE case_id=?", (case_id,)).fetchone()[0],
            1,
        )

    def test_child_records_and_task_actions_are_isolated_between_firms(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        created = self._mutate("post", "/cases", json={"title": "Firm A child-record matter"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        task_response = self._mutate(
            "post", f"/cases/{case_id}/tasks", json={"title": "Firm A confidential task"}
        )
        self.assertEqual(task_response.status_code, 200, task_response.text)

        context_a = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_a_db = self.runtime.firm_connection(context_a)
        task = firm_a_db.execute(
            "SELECT id,title FROM tasks WHERE case_id=?", (case_id,)
        ).fetchone()
        self.assertIsNotNone(task)
        task_id = task["id"]

        switched = self._mutate(
            "post", "/session/firm", json={"firm_id": self.firm_b}
        )
        self.assertEqual(switched.status_code, 200, switched.text)
        self.assertEqual(self.client.get("/cases").json(), [])

        # A child-record ID from Firm A must not resolve against Firm B's DB.
        cross_firm_action = self._mutate("post", f"/tasks/{task_id}/done")
        self.assertEqual(cross_firm_action.status_code, 404, cross_firm_action.text)
        context_b = self.runtime.resolve_context(self.user_id, self.firm_b)
        firm_b_db = self.runtime.firm_connection(context_b)
        self.assertEqual(
            firm_b_db.execute("SELECT COUNT(*) FROM tasks").fetchone()[0], 0
        )
        self.assertEqual(
            firm_a_db.execute("SELECT title FROM tasks WHERE id=?", (task_id,)).fetchone()[0],
            "Firm A confidential task",
        )

    def test_missing_tenant_catalog_entry_fails_closed_without_fallback(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "Missing Catalog Guard"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        control = self.runtime.control_connection()
        catalog = control.execute(
            "SELECT database_file,status FROM tenant_database_catalog WHERE firm_id=?",
            (self.firm_a,),
        ).fetchone()
        self.assertIsNotNone(catalog)
        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        control.execute("DELETE FROM tenant_database_catalog WHERE firm_id=?", (self.firm_a,))
        control.commit()
        try:
            detail = self.client.get(f"/cases/{case_id}")
            self.assertEqual(detail.status_code, 503, detail.text)
            create_attempt = self._mutate(
                "post", "/cases", json={"title": "Must not bypass missing catalog"}
            )
            self.assertEqual(create_attempt.status_code, 503, create_attempt.text)
            self.assertEqual(
                firm_db.execute("SELECT COUNT(*) FROM cases").fetchone()[0],
                1,
                "missing catalog entries must not trigger shared-database fallback writes",
            )
        finally:
            control.execute(
                "INSERT INTO tenant_database_catalog(firm_id,database_file,status) VALUES(?,?,?)",
                (self.firm_a, catalog["database_file"], catalog["status"]),
            )
            control.commit()

        recovered = self.client.get(f"/cases/{case_id}")
        self.assertEqual(recovered.status_code, 200, recovered.text)
        self.assertEqual(recovered.json()["id"], case_id)

    def test_unavailable_tenant_database_fails_closed_without_fallback(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "Unavailable Tenant Guard"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        control = self.runtime.control_connection()
        catalog = control.execute(
            "SELECT database_file FROM tenant_database_catalog WHERE firm_id=?",
            (self.firm_a,),
        ).fetchone()
        self.assertIsNotNone(catalog)
        tenant_path = self.root / "tenants" / catalog["database_file"]
        unavailable_path = tenant_path.with_name(tenant_path.name + ".temporarily-unavailable")
        self.assertTrue(tenant_path.is_file())
        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        os.replace(tenant_path, unavailable_path)
        try:
            detail = self.client.get(f"/cases/{case_id}")
            self.assertEqual(detail.status_code, 503, detail.text)

            create_attempt = self._mutate(
                "post", "/cases", json={"title": "Must not fall back to another database"}
            )
            self.assertEqual(create_attempt.status_code, 503, create_attempt.text)

            # The still-open trusted connection sees the original tenant inode.
            # No write may have reached it while catalog resolution was failing.
            self.assertEqual(
                firm_db.execute("SELECT COUNT(*) FROM cases").fetchone()[0],
                1,
                "unavailable tenant reads/writes must fail before any fallback write",
            )
        finally:
            os.replace(unavailable_path, tenant_path)

        recovered = self.client.get(f"/cases/{case_id}")
        self.assertEqual(recovered.status_code, 200, recovered.text)
        self.assertEqual(recovered.json()["id"], case_id)

    def test_disabled_firm_is_denied_at_api_and_retains_its_database(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        created = self._mutate(
            "post", "/cases", json={"title": "Disabled firm retained sentinel"}
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        control = self.runtime.control_connection()
        filename = control.execute(
            "SELECT database_file FROM tenant_database_catalog WHERE firm_id=?",
            (self.firm_b,),
        ).fetchone()["database_file"]
        database_path = self.root / "tenants" / filename

        self.assertTrue(self.runtime.disable_firm(self.firm_b))
        denied = self.client.get("/cases")
        self.assertEqual(denied.status_code, 403, denied.text)
        status = control.execute("SELECT status FROM firms WHERE id=?", (self.firm_b,)).fetchone()["status"]
        self.assertEqual(status, "disabled")
        self.assertTrue(database_path.is_file())

        import sqlite3
        read_only = sqlite3.connect(database_path.resolve().as_uri() + "?mode=ro", uri=True)
        try:
            row = read_only.execute("SELECT title FROM cases WHERE id=?", (case_id,)).fetchone()
            self.assertEqual(row[0], "Disabled firm retained sentinel")
        finally:
            read_only.close()

    def test_password_change_audit_failure_rolls_back_credential_and_sessions(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        control = self.runtime.control_connection()
        control.execute("""
            CREATE TRIGGER reject_web_password_audit BEFORE INSERT ON platform_security_events
            WHEN NEW.event_type='password_changed_web'
            BEGIN SELECT RAISE(ABORT, 'password audit unavailable'); END
        """)
        control.commit()
        before_hash = control.execute("SELECT password_hash FROM users WHERE id=?", (self.user_id,)).fetchone()[0]
        response = self._mutate(
            "post", "/account/password",
            json={"current_password": self.password,
                  "new_password": "EvenLongerUniqueTestPassphrase-2026!"},
        )
        self.assertEqual(response.status_code, 503, response.text)
        control = self.runtime.control_connection()
        after_hash = control.execute("SELECT password_hash FROM users WHERE id=?", (self.user_id,)).fetchone()[0]
        self.assertEqual(after_hash, before_hash)
        self.assertGreater(control.execute("SELECT COUNT(*) FROM auth_sessions WHERE user_id=?", (self.user_id,)).fetchone()[0], 0)
        self.assertEqual(control.execute(
            "SELECT COUNT(*) FROM platform_security_events WHERE event_type='password_changed_web' AND target_user_id=?",
            (self.user_id,),
        ).fetchone()[0], 0)

    def test_password_change_revokes_sessions_across_all_firm_memberships(self):
        first_login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(first_login.status_code, 200, first_login.text)
        from fastapi.testclient import TestClient
        other_client = TestClient(app.app)
        try:
            second_login = other_client.post(
                "/login", json={"email": "OWNER@example.test", "password": self.password}
            )
            self.assertEqual(second_login.status_code, 200, second_login.text)
            switch = other_client.post(
                "/session/firm", json={"firm_id": self.firm_b},
                headers={"X-CSRF-Token": other_client.cookies.get("lc_csrf") or ""},
            )
            self.assertEqual(switch.status_code, 200, switch.text)
            self.assertEqual(other_client.get("/cases").status_code, 200)

            changed = self._mutate(
                "post", "/account/password",
                json={"current_password": self.password,
                      "new_password": "EvenLongerUniqueTestPassphrase-2026!"},
            )
            self.assertEqual(changed.status_code, 200, changed.text)
            self.assertEqual(other_client.get("/cases").status_code, 401)
            self.assertEqual(
                self.runtime.control_connection().execute(
                    "SELECT COUNT(*) FROM auth_sessions WHERE user_id=?", (self.user_id,)
                ).fetchone()[0],
                0,
            )
            event = self.runtime.control_connection().execute(
                "SELECT actor_user_id,target_user_id,event_type,detail_json FROM platform_security_events "
                "WHERE event_type='password_changed_web' AND target_user_id=? ORDER BY id DESC LIMIT 1",
                (self.user_id,),
            ).fetchone()
            self.assertIsNotNone(event)
            self.assertEqual(event["actor_user_id"], self.user_id)
            self.assertIn('"sessions_revoked": true', event["detail_json"])
            self.assertEqual(
                self.runtime.control_connection().execute(
                    "SELECT status FROM firm_memberships WHERE user_id=? AND firm_id=?",
                    (self.user_id, self.firm_b),
                ).fetchone()["status"],
                "active",
            )
            new_login = self.client.post(
                "/login", json={"email": "OWNER@example.test",
                                "password": "EvenLongerUniqueTestPassphrase-2026!"},
            )
            self.assertEqual(new_login.status_code, 200, new_login.text)
        finally:
            other_client.close()

    def test_password_change_rejects_wrong_current_short_and_reused_passwords(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        control = self.runtime.control_connection()
        before = control.execute(
            "SELECT password_hash FROM users WHERE id=?", (self.user_id,)
        ).fetchone()["password_hash"]

        wrong = self._mutate(
            "post", "/account/password",
            json={"current_password": "incorrect-current-password",
                  "new_password": "ReplacementPasswordPassphrase-2026!"},
        )
        self.assertEqual(wrong.status_code, 403, wrong.text)
        short = self._mutate(
            "post", "/account/password",
            json={"current_password": self.password, "new_password": "short"},
        )
        self.assertEqual(short.status_code, 422, short.text)
        reused = self._mutate(
            "post", "/account/password",
            json={"current_password": self.password, "new_password": self.password},
        )
        self.assertEqual(reused.status_code, 422, reused.text)

        after = control.execute(
            "SELECT password_hash FROM users WHERE id=?", (self.user_id,)
        ).fetchone()["password_hash"]
        self.assertEqual(after, before)
        self.assertEqual(self.client.get("/session").status_code, 200)

    def test_malformed_authorization_missing_csrf_and_expired_session_fail_closed(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        with TestClient(app.app) as unauthenticated:
            for malformed in ("", "garbage", "Bearer", "Basic abc", "Bearer not-a-real-token"):
                response = unauthenticated.get("/cases", headers={"Authorization": malformed})
                self.assertEqual(response.status_code, 401, malformed)

        missing_csrf = self.client.post("/cases", json={"title": "Missing CSRF"})
        self.assertEqual(missing_csrf.status_code, 403, missing_csrf.text)

        token = self.client.cookies.get("lc_session") or ""
        self.assertTrue(token)
        control = self.runtime.control_connection()
        control.execute(
            "UPDATE auth_sessions SET expires=? WHERE token_hash=?",
            (int(app.time.time()) - 1, app.sha(token)),
        )
        control.commit()
        expired = self.client.get("/cases")
        self.assertEqual(expired.status_code, 401, expired.text)

    def test_deadline_action_cannot_mutate_another_firms_deadline(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "Firm A deadline action matter"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        created_deadline = self._mutate(
            "post", f"/cases/{case_id}/deadlines",
            json={"title": "Firm A protected deadline", "due": "2099-01-01"},
        )
        self.assertEqual(created_deadline.status_code, 200, created_deadline.text)
        context_a = self.runtime.resolve_context(self.user_id, self.firm_a)
        db_a = self.runtime.firm_connection(context_a)
        deadline_id = db_a.execute(
            "SELECT id FROM deadlines WHERE case_id=?", (case_id,)
        ).fetchone()["id"]

        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        denied = self._mutate("post", f"/deadlines/{deadline_id}/done")
        self.assertEqual(denied.status_code, 404, denied.text)
        self.assertEqual(
            db_a.execute("SELECT status FROM deadlines WHERE id=?", (deadline_id,)).fetchone()["status"],
            "open",
        )
        context_b = self.runtime.resolve_context(self.user_id, self.firm_b)
        db_b = self.runtime.firm_connection(context_b)
        self.assertEqual(db_b.execute("SELECT COUNT(*) FROM deadlines").fetchone()[0], 0)

    def test_cross_firm_task_claim_and_completion_are_denied(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "Firm A task action matter"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        created_task = self._mutate(
            "post", f"/cases/{case_id}/tasks", json={"title": "Firm A protected task"}
        )
        self.assertEqual(created_task.status_code, 200, created_task.text)
        context_a = self.runtime.resolve_context(self.user_id, self.firm_a)
        db_a = self.runtime.firm_connection(context_a)
        task = db_a.execute(
            "SELECT id,status,assignee FROM tasks WHERE case_id=?", (case_id,)
        ).fetchone()
        task_id = task["id"]

        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        self.assertEqual(self._mutate("post", f"/tasks/{task_id}/claim").status_code, 404)
        self.assertEqual(self._mutate("post", f"/tasks/{task_id}/done").status_code, 404)
        after = db_a.execute(
            "SELECT status,assignee FROM tasks WHERE id=?", (task_id,)
        ).fetchone()
        self.assertEqual(after["status"], "open")
        self.assertIsNone(after["assignee"])
        context_b = self.runtime.resolve_context(self.user_id, self.firm_b)
        db_b = self.runtime.firm_connection(context_b)
        self.assertEqual(db_b.execute("SELECT COUNT(*) FROM tasks").fetchone()[0], 0)

    def test_nonmember_cannot_switch_firms_and_revoked_membership_blocks_cached_session(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        unassigned_firm, _ = self.runtime.provision_firm("Unassigned Firm")
        denied = self._mutate(
            "post", "/session/firm", json={"firm_id": unassigned_firm}
        )
        self.assertEqual(denied.status_code, 403, denied.text)

        switched = self._mutate(
            "post", "/session/firm", json={"firm_id": self.firm_b}
        )
        self.assertEqual(switched.status_code, 200, switched.text)
        # Prime the tenant connection cache and create case-scoped resources before
        # revocation so direct resource routes are tested with a stale session too.
        self.assertEqual(self.client.get("/cases").status_code, 200)
        created_case = self._mutate(
            "post", "/cases", json={"title": "REVOKED_MEMBERSHIP_CASE_SENTINEL"}
        )
        self.assertEqual(created_case.status_code, 200, created_case.text)
        case_id = created_case.json()["id"]
        created_document = self._mutate(
            "post", f"/cases/{case_id}/docs",
            json={"name": "revoked-membership.txt", "text": "REVOKED_MEMBERSHIP_DOCUMENT_SENTINEL"},
        )
        self.assertEqual(created_document.status_code, 200, created_document.text)
        document_id = created_document.json()["document_id"]

        control = self.runtime.control_connection()
        control.execute(
            "UPDATE firm_memberships SET status='revoked', updated_at=datetime('now') "
            "WHERE user_id=? AND firm_id=?",
            (self.user_id, self.firm_b),
        )
        control.commit()

        # Every request must revalidate membership, even with a cached firm DB
        # connection and a still-valid authentication session cookie.
        denied_after_revocation = self.client.get("/cases")
        self.assertEqual(denied_after_revocation.status_code, 403,
                         denied_after_revocation.text)

        # The same stale cookie must not retain access through other route
        # families after the control-plane membership is revoked.
        identity_before = control.execute(
            "SELECT password_hash FROM users WHERE id=?", (self.user_id,)
        ).fetchone()["password_hash"]
        mfa_event_count_before = control.execute(
            "SELECT COUNT(*) FROM mfa_security_events WHERE actor_user_id=?", (self.user_id,)
        ).fetchone()[0]
        revoked_route_attempts = [
            ("get", "/session", None),
            ("get", "/dashboard", None),
            ("get", "/calendar.ics", None),
            ("get", "/finance/overview", None),
            ("get", "/research/conversations", None),
            ("get", "/invitations", None),
            ("get", f"/cases/{case_id}", None),
            ("get", f"/cases/{case_id}/docs/{document_id}/reviews", None),
            ("get", f"/cases/{case_id}/docs/{document_id}/original", None),
            ("get", f"/cases/{case_id}/docs/{document_id}/original/legal-hold/events", None),
            ("get", f"/cases/{case_id}/docs/{document_id}/original/deletion-requests", None),
            ("post", f"/cases/{case_id}/tasks", {"title": "must be denied"}),
            ("get", "/audit/case-deletions", None),
            ("get", "/audit/security-events", None),
            ("get", "/audit/security-events.csv", None),
            ("post", "/search", {"q": "FIRM_B_PRIVATE_SENTINEL"}),
            ("post", "/ask", {"q": "FIRM_B_PRIVATE_SENTINEL"}),
            # Revocation must also block identity/account-security operations,
            # not merely matter and research endpoints.
            ("post", "/account/password", {
                "current_password": self.password,
                "new_password": "RevokedAccountNewPassword-2026!",
            }),
            ("post", "/account/mfa/enroll", {"current_password": self.password}),
            ("post", "/account/mfa/confirm", {
                "current_password": self.password, "code": "123456",
            }),
            ("post", "/account/mfa/verify", {"code": "123456"}),
            ("post", "/account/mfa/recovery-codes/rotate", {
                "current_password": self.password, "code": "123456",
            }),
            ("post", "/account/mfa/re-enroll", {"current_password": self.password}),
            ("post", "/account/mfa/disable", {
                "current_password": self.password, "code": "123456",
            }),
        ]
        for method, path, payload in revoked_route_attempts:
            with self.subTest(method=method, path=path):
                kwargs = {}
                if payload is not None:
                    kwargs["json"] = payload
                if method == "post":
                    kwargs["headers"] = {"X-CSRF-Token": self.client.cookies.get("lc_csrf") or ""}
                response = getattr(self.client, method)(path, **kwargs)
                self.assertEqual(response.status_code, 403, response.text)

        self.assertEqual(
            self.runtime.firm_connection(
                self.runtime.resolve_context(self.user_id, self.firm_a)
            ).execute("SELECT COUNT(*) FROM cases").fetchone()[0],
            0,
        )
        self.assertEqual(
            control.execute("SELECT password_hash FROM users WHERE id=?", (self.user_id,)).fetchone()["password_hash"],
            identity_before,
        )
        self.assertIsNone(control.execute(
            "SELECT user_id FROM mfa_credentials WHERE user_id=?", (self.user_id,)
        ).fetchone())
        self.assertEqual(
            control.execute(
                "SELECT COUNT(*) FROM mfa_security_events WHERE actor_user_id=?", (self.user_id,)
            ).fetchone()[0],
            mfa_event_count_before,
        )

    def test_audit_lists_and_csv_exports_are_scoped_to_active_firm(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        context_a = self.runtime.resolve_context(self.user_id, self.firm_a)
        db_a = self.runtime.firm_connection(context_a)
        db_a.execute(
            "INSERT INTO case_deletion_audit(firm_id,case_id,actor_user_id,actor_name,"
            "case_snapshot_json,events_snapshot_json,related_counts_json) VALUES(?,?,?,?,?,?,?)",
            (self.firm_a, 101, self.user_id, "Owner A", '{"title":"FIRM_A_AUDIT_SENTINEL"}', "[]", "{}"),
        )
        db_a.execute(
            "INSERT INTO security_audit_events(firm_id,user_id,event) VALUES(?,?,?)",
            (self.firm_a, self.user_id, "FIRM_A_SECURITY_SENTINEL"),
        )
        db_a.commit()

        context_b = self.runtime.resolve_context(self.user_id, self.firm_b)
        db_b = self.runtime.firm_connection(context_b)
        db_b.execute(
            "INSERT INTO case_deletion_audit(firm_id,case_id,actor_user_id,actor_name,"
            "case_snapshot_json,events_snapshot_json,related_counts_json) VALUES(?,?,?,?,?,?,?)",
            (self.firm_b, 202, self.user_id, "Owner B", '{"title":"FIRM_B_AUDIT_SENTINEL"}', "[]", "{}"),
        )
        db_b.execute(
            "INSERT INTO security_audit_events(firm_id,user_id,event) VALUES(?,?,?)",
            (self.firm_b, self.user_id, "FIRM_B_SECURITY_SENTINEL"),
        )
        db_b.commit()

        audit_a = self.client.get("/audit/case-deletions")
        self.assertEqual(audit_a.status_code, 200, audit_a.text)
        self.assertEqual([item["firm_id"] for item in audit_a.json()["items"]], [self.firm_a])
        self.assertIn("FIRM_A_AUDIT_SENTINEL", audit_a.text)
        self.assertNotIn("FIRM_B_AUDIT_SENTINEL", audit_a.text)

        events_a = self.client.get("/audit/security-events")
        self.assertEqual(events_a.status_code, 200, events_a.text)
        self.assertEqual([item["firm_id"] for item in events_a.json()["items"]], [self.firm_a])
        csv_a = self.client.get("/audit/security-events.csv")
        self.assertEqual(csv_a.status_code, 200, csv_a.text)
        self.assertIn("FIRM_A_SECURITY_SENTINEL", csv_a.text)
        self.assertNotIn("FIRM_B_SECURITY_SENTINEL", csv_a.text)

        control = self.runtime.control_connection()
        control.execute(
            "UPDATE firm_memberships SET firm_role='admin' WHERE user_id=? AND firm_id=?",
            (self.user_id, self.firm_b),
        )
        control.commit()
        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)

        audit_b = self.client.get("/audit/case-deletions")
        self.assertEqual(audit_b.status_code, 200, audit_b.text)
        self.assertEqual([item["firm_id"] for item in audit_b.json()["items"]], [self.firm_b])
        self.assertIn("FIRM_B_AUDIT_SENTINEL", audit_b.text)
        self.assertNotIn("FIRM_A_AUDIT_SENTINEL", audit_b.text)

        events_b = self.client.get("/audit/security-events")
        self.assertEqual(events_b.status_code, 200, events_b.text)
        self.assertEqual([item["firm_id"] for item in events_b.json()["items"]], [self.firm_b])
        csv_b = self.client.get("/audit/security-events.csv")
        self.assertEqual(csv_b.status_code, 200, csv_b.text)
        self.assertIn("FIRM_B_SECURITY_SENTINEL", csv_b.text)
        self.assertNotIn("FIRM_A_SECURITY_SENTINEL", csv_b.text)

    def test_document_extraction_review_routes_are_isolated_across_firms(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        created_case = self._mutate("post", "/cases", json={"title": "Firm A review matter"})
        self.assertEqual(created_case.status_code, 200, created_case.text)
        case_id = created_case.json()["id"]
        uploaded = self._mutate(
            "post",
            f"/cases/{case_id}/docs",
            json={"name": "private-review.txt", "text": "CONFIDENTIAL-REVIEW-SENTINEL " * 12},
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        document_id = uploaded.json()["document_id"]
        context_a = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_a_db = self.runtime.firm_connection(context_a)
        chunk = firm_a_db.execute(
            "SELECT id,text FROM chunks WHERE document_id=? AND case_id=? ORDER BY id LIMIT 1",
            (document_id, case_id),
        ).fetchone()
        self.assertIsNotNone(chunk)
        chunk_id, original_text = chunk["id"], chunk["text"]
        import hashlib
        original_digest = hashlib.sha256(original_text.encode("utf-8")).hexdigest()

        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        for method, url, kwargs in (
            ("get", f"/cases/{case_id}/docs/{document_id}/reviews", {}),
            ("get", f"/cases/{case_id}/docs/{document_id}/chunks", {}),
            ("post", f"/cases/{case_id}/docs/{document_id}/chunks/{chunk_id}/review", {
                "json": {
                    "outcome": "corrected",
                    "corrected_text": "UNAUTHORIZED-CROSS-FIRM-CORRECTION",
                    "expected_text_sha256": original_digest,
                },
            }),
        ):
            response = self._mutate(method, url, **kwargs) if method != "get" else self.client.get(url)
            self.assertEqual(response.status_code, 404, response.text)

        # Denied review requests must not alter the source firm's extracted text.
        unchanged = firm_a_db.execute(
            "SELECT text FROM chunks WHERE id=? AND document_id=? AND case_id=?",
            (chunk_id, document_id, case_id),
        ).fetchone()
        self.assertIsNotNone(unchanged)
        self.assertEqual(unchanged["text"], original_text)
        self.assertEqual(
            firm_a_db.execute(
                "SELECT COUNT(*) FROM document_extraction_reviews WHERE chunk_id=?",
                (chunk_id,),
            ).fetchone()[0],
            0,
        )

    def test_document_review_routes_require_document_to_match_path_case(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        case_one = self._mutate("post", "/cases", json={"title": "Document owner case"})
        case_two = self._mutate("post", "/cases", json={"title": "Different path case"})
        self.assertEqual(case_one.status_code, 200, case_one.text)
        self.assertEqual(case_two.status_code, 200, case_two.text)
        case_one_id, case_two_id = case_one.json()["id"], case_two.json()["id"]
        uploaded = self._mutate(
            "post", f"/cases/{case_one_id}/docs",
            json={"name": "bound-document.txt", "text": "CASE_DOCUMENT_BINDING_SENTINEL " * 12},
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        document_id = uploaded.json()["document_id"]

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        chunk = firm_db.execute(
            "SELECT id,text FROM chunks WHERE document_id=? AND case_id=? ORDER BY id LIMIT 1",
            (document_id, case_one_id),
        ).fetchone()
        self.assertIsNotNone(chunk)
        chunk_id, original_text = int(chunk["id"]), str(chunk["text"])
        import hashlib
        digest = hashlib.sha256(original_text.encode("utf-8")).hexdigest()

        # The user is authorized for both cases. A caller must still not be able
        # to pair a document/chunk ID from one case with another case's URL.
        for method, path, kwargs in (
            ("get", f"/cases/{case_two_id}/docs/{document_id}/reviews", {}),
            ("get", f"/cases/{case_two_id}/docs/{document_id}/chunks", {}),
            ("post", f"/cases/{case_two_id}/docs/{document_id}/chunks/{chunk_id}/review", {
                "json": {
                    "outcome": "corrected",
                    "corrected_text": "UNAUTHORIZED-CROSS-CASE-CORRECTION",
                    "expected_text_sha256": digest,
                },
            }),
        ):
            response = self._mutate(method, path, **kwargs) if method != "get" else self.client.get(path)
            self.assertEqual(response.status_code, 404, f"{path}: {response.text}")

        self.assertEqual(
            firm_db.execute(
                "SELECT text FROM chunks WHERE id=? AND document_id=? AND case_id=?",
                (chunk_id, document_id, case_one_id),
            ).fetchone()["text"],
            original_text,
        )
        self.assertEqual(
            firm_db.execute(
                "SELECT COUNT(*) FROM document_extraction_reviews WHERE chunk_id=?",
                (chunk_id,),
            ).fetchone()[0],
            0,
        )

    def test_same_firm_nonmember_cannot_read_or_mutate_extraction_reviews(self):
        owner_login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(owner_login.status_code, 200, owner_login.text)
        created_case = self._mutate("post", "/cases", json={"title": "Case-member review boundary"})
        self.assertEqual(created_case.status_code, 200, created_case.text)
        case_id = created_case.json()["id"]
        uploaded = self._mutate(
            "post", f"/cases/{case_id}/docs",
            json={"name": "private-extraction.txt", "text": "SAME_FIRM_PRIVATE_EXTRACTION " * 12},
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        document_id = uploaded.json()["document_id"]

        owner_context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(owner_context)
        chunk = firm_db.execute(
            "SELECT id,text FROM chunks WHERE document_id=? AND case_id=? ORDER BY id LIMIT 1",
            (document_id, case_id),
        ).fetchone()
        self.assertIsNotNone(chunk)
        chunk_id, original_text = int(chunk["id"]), str(chunk["text"])
        import hashlib
        digest = hashlib.sha256(original_text.encode("utf-8")).hexdigest()

        other_password = "SameFirmNonmember-Password-2026!"
        other_user_id = self.runtime.create_global_identity(
            "Same Firm Nonmember", "nonmember-review@example.test", app.hash_pw(other_password)
        )
        self.runtime.create_firm_membership(other_user_id, self.firm_a, "lawyer")

        other_login = self.client.post(
            "/login",
            json={"email": "nonmember-review@example.test", "password": other_password},
        )
        self.assertEqual(other_login.status_code, 200, other_login.text)
        for method, path, kwargs in (
            ("get", f"/cases/{case_id}/docs/{document_id}/reviews", {}),
            ("get", f"/cases/{case_id}/docs/{document_id}/chunks", {}),
            ("post", f"/cases/{case_id}/docs/{document_id}/chunks/{chunk_id}/review", {
                "json": {
                    "outcome": "corrected",
                    "corrected_text": "UNAUTHORIZED-SAME-FIRM-CORRECTION",
                    "expected_text_sha256": digest,
                },
            }),
        ):
            response = self._mutate(method, path, **kwargs) if method != "get" else self.client.get(path)
            self.assertEqual(response.status_code, 404, response.text)

        unchanged = firm_db.execute(
            "SELECT text FROM chunks WHERE id=? AND document_id=? AND case_id=?",
            (chunk_id, document_id, case_id),
        ).fetchone()
        self.assertIsNotNone(unchanged)
        self.assertEqual(unchanged["text"], original_text)
        self.assertEqual(
            firm_db.execute(
                "SELECT COUNT(*) FROM document_extraction_reviews WHERE chunk_id=?",
                (chunk_id,),
            ).fetchone()[0],
            0,
        )

    def test_research_conversations_are_isolated_across_firms(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        created = self._mutate("post", "/research/conversations")
        self.assertEqual(created.status_code, 200, created.text)
        conversation_id = created.json()["id"]
        context_a = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_a_db = self.runtime.firm_connection(context_a)
        self.assertEqual(
            firm_a_db.execute(
                "SELECT COUNT(*) FROM research_conversations WHERE id=? AND firm_id=? AND user_id=?",
                (conversation_id, self.firm_a, self.user_id),
            ).fetchone()[0],
            1,
        )

        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        self.assertEqual(self.client.get("/research/conversations").json(), [])

        detail = self.client.get(f"/research/conversations/{conversation_id}")
        self.assertEqual(detail.status_code, 404, detail.text)
        rename = self._mutate(
            "patch", f"/research/conversations/{conversation_id}",
            json={"title": "Unauthorized cross-firm rename"},
        )
        self.assertEqual(rename.status_code, 404, rename.text)
        send_message = self._mutate(
            "post", f"/research/conversations/{conversation_id}/messages",
            json={"message": "CONFIDENTIAL-RESEARCH-SENTINEL", "scope": "verified_corpus"},
        )
        self.assertEqual(send_message.status_code, 404, send_message.text)
        save_to_matter = self._mutate(
            "post", f"/research/conversations/{conversation_id}/save-to-matter",
            json={"case_id": 1, "message_id": 1},
        )
        self.assertEqual(save_to_matter.status_code, 404, save_to_matter.text)
        delete = self._mutate("delete", f"/research/conversations/{conversation_id}")
        self.assertEqual(delete.status_code, 404, delete.text)

        # Attempts in Firm B must not mutate or delete Firm A's conversation.
        self.assertEqual(
            firm_a_db.execute(
                "SELECT COUNT(*) FROM research_conversations WHERE id=? AND firm_id=?",
                (conversation_id, self.firm_a),
            ).fetchone()[0],
            1,
        )
        self.assertEqual(
            firm_a_db.execute(
                "SELECT COUNT(*) FROM research_messages WHERE conversation_id=?",
                (conversation_id,),
            ).fetchone()[0],
            0,
        )


    def test_research_answers_receive_stable_opaque_ids_in_send_and_history(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/research/conversations")
        self.assertEqual(created.status_code, 200, created.text)
        conversation_id = created.json()["id"]

        synthetic_result = {
            "found": True,
            "answer": "Synthetic answer [1].",
            "abstained": False,
            "sources": [{"ref": "SYNTHETIC-LAW-1", "text": "Synthetic source passage."}],
            "scope": "verified_corpus",
            "case_id": None,
            "requested_as_of": "2026-11-01",
        }
        with patch("app.ask", return_value=synthetic_result):
            first = self._mutate(
                "post",
                f"/research/conversations/{conversation_id}/messages",
                json={"message": "Synthetic question", "scope": "verified_corpus"},
            )
            second = self._mutate(
                "post",
                f"/research/conversations/{conversation_id}/messages",
                json={"message": "Synthetic question", "scope": "verified_corpus"},
            )

        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(second.status_code, 200, second.text)
        first_id = first.json()["answer_id"]
        second_id = second.json()["answer_id"]
        self.assertRegex(first_id, r"^ans_[A-Za-z0-9_-]{32}$")
        self.assertRegex(second_id, r"^ans_[A-Za-z0-9_-]{32}$")
        self.assertNotEqual(first_id, second_id, "each generated answer needs its own identifier")

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        # Historical assistant rows without an answer_id remain readable and unmodified.
        firm_db.execute(
            "INSERT INTO research_messages(conversation_id,role,content,sources_json,message_json) "
            "VALUES(?,?,?,?,?)",
            (conversation_id, "assistant", "Legacy synthetic answer", "[]",
             json.dumps({"scope": "verified_corpus", "case_id": None})),
        )
        firm_db.commit()

        detail = self.client.get(f"/research/conversations/{conversation_id}")
        self.assertEqual(detail.status_code, 200, detail.text)
        assistant_messages = [
            item for item in detail.json()["messages"] if item["role"] == "assistant"
        ]
        self.assertEqual(len(assistant_messages), 3)
        self.assertEqual(
            [item["answer_id"] for item in assistant_messages[:2]],
            [first_id, second_id],
        )
        self.assertEqual(
            [item["result"]["answer_id"] for item in assistant_messages[:2]],
            [first_id, second_id],
        )
        self.assertNotIn("answer_id", assistant_messages[2])
        self.assertNotIn("answer_id", assistant_messages[2]["result"])

        stored = firm_db.execute(
            "SELECT message_json FROM research_messages "
            "WHERE conversation_id=? AND role='assistant' ORDER BY id",
            (conversation_id,),
        ).fetchall()
        self.assertEqual(
            [json.loads(row["message_json"]).get("answer_id") for row in stored],
            [first_id, second_id, None],
        )

    def test_login_firm_switch_and_shared_corpus_are_separated(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        self.assertEqual(len(login.json()["firms"]), 2)
        self.assertEqual(login.json()["firm_id"], self.firm_a)

        created = self._mutate("post", "/cases", json={"title": "Firm A matter"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        uploaded = self._mutate(
            "post",
            f"/cases/{case_id}/docs",
            json={"name": "private.txt", "text": "firmonlysentinel confidential matter passage"},
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        matter_search = self._mutate(
            "post",
            "/search",
            json={"q": "firmonlysentinel", "case_id": case_id, "scope": "firm_documents"},
        )
        self.assertEqual(matter_search.status_code, 200, matter_search.text)
        self.assertTrue(any("firmonlysentinel" in hit["text"] for hit in matter_search.json()))

        # The legal corpus is central and shared; matter content is not.
        corpus = self.runtime.legal_corpus_connection()
        corpus.execute(
            "INSERT INTO legal_sources(source_id,authority,source_page,collection_permission,"
            "commercial_reuse_permission,permission_evidence) VALUES(?,?,?,?,?,?)",
            ("test-source", "Test authority", "https://example.test/law", "allowed", "allowed", "test evidence"),
        )
        corpus.execute(
            "INSERT INTO legal_documents(document_id,source_id,official_title,document_type,practice_area) "
            "VALUES(?,?,?,?,?)",
            ("test-law", "test-source", "Test Law", "statute", "general"),
        )
        version_id = corpus.execute(
            "INSERT INTO legal_document_versions(document_id,version_label,source_url,effective_date,"
            "legal_status,review_status,authenticity_verified,status_verified,amendment_history_verified,"
            "permission_evidence) VALUES(?,?,?,?,?,?,?,?,?,?)",
            ("test-law", "current", "https://example.test/law/current", "2020-01-01",
             "in_force", "verified", 1, 1, 1, "version evidence"),
        ).lastrowid
        for review_type in ("source_authenticity", "extraction_quality", "legal_status", "amendment_history", "licensing", "lawyer_approval"):
            notes = "Reviewed test fixture."
            if review_type == "lawyer_approval":
                notes = "[qualified_egyptian_counsel_attested:user_id=42] " + notes
            corpus.execute(
                "INSERT INTO legal_review_events(version_id,reviewer,review_type,outcome,notes,evidence_url) VALUES(?,?,?,?,?,?)",
                (version_id, "Test reviewer", review_type, "pass", notes, "https://example.test/review"),
            )
        corpus.execute(
            "INSERT INTO chunks(ref,article,page,text,norm,document_version_id,in_force,as_of,source,verified) "
            "VALUES(?,?,?,?,?,?,?,?,?,1)",
            ("Test Law Article 1", "1", "1", "corpusonlysentinel statutory rule",
             "corpusonlysentinel statutory rule", version_id, 1, "2020-01-01", "test-source"),
        )
        corpus.commit()

        # A firm admin is not a platform legal reviewer.
        self.assertEqual(self.client.get("/legal-review/queue").status_code, 403)

        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        self.assertEqual(switched.json()["firm_id"], self.firm_b)
        self.assertEqual(self.client.get("/cases").json(), [])
        guessed_case = self._mutate(
            "post",
            "/search",
            json={"q": "firmonlysentinel", "case_id": case_id, "scope": "firm_documents"},
        )
        self.assertEqual(guessed_case.status_code, 404)

        shared_search = self._mutate(
            "post", "/search", json={"q": "corpusonlysentinel", "scope": "verified_corpus"}
        )
        self.assertEqual(shared_search.status_code, 200, shared_search.text)
        self.assertTrue(any("corpusonlysentinel" in hit["text"] for hit in shared_search.json()))

        control = self.runtime.control_connection()
        control.execute(
            "INSERT INTO user_platform_roles(user_id,role_name) VALUES(?,'platform_reviewer')",
            (self.user_id,),
        )
        control.commit()
        reviewer_queue = self.client.get("/legal-review/queue")
        self.assertEqual(reviewer_queue.status_code, 200, reviewer_queue.text)


    def test_shared_corpus_import_requires_platform_reviewer_not_firm_admin(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        payload = [{
            "law": "Synthetic staged law import",
            "article": "1",
            "text": "Synthetic passage for platform authorization regression coverage.",
            "source": "TEST-ONLY",
            "verified": 1,
            "in_force": 1,
        }]
        headers = {"X-Admin-Token": self.admin_token}

        # The initial owner is a firm administrator, not a platform reviewer.
        denied = self._mutate("post", "/law/import", json=payload, headers=headers)
        self.assertEqual(denied.status_code, 403, denied.text)

        # The retired shared-database mode must not restore the former admin/token bypass.
        with patch.object(app, "_separate_runtime_enabled", return_value=False), patch.object(
            app, "control_conn", return_value=self.runtime.control_connection()
        ):
            with self.assertRaises(HTTPException) as legacy_denied:
                app.law_import(
                    [app.LawRow(**payload[0])],
                    self.admin_token,
                    {"id": self.user_id, "role": "admin", "email": "owner@example.test"},
                )
        self.assertEqual(legacy_denied.exception.status_code, 403)
        corpus = self.runtime.legal_corpus_connection()
        self.assertEqual(
            corpus.execute(
                "SELECT COUNT(*) FROM legal_import_staging WHERE law=?",
                ("Synthetic staged law import",),
            ).fetchone()[0],
            0,
        )

        control = self.runtime.control_connection()
        control.execute(
            "INSERT INTO user_platform_roles(user_id,role_name) "
            "VALUES(?,'platform_reviewer')",
            (self.user_id,),
        )
        control.commit()
        allowed = self._mutate("post", "/law/import", json=payload, headers=headers)
        self.assertEqual(allowed.status_code, 200, allowed.text)
        self.assertTrue(allowed.json()["review_required"])
        self.assertFalse(allowed.json()["search_eligible"])
        self.assertEqual(
            corpus.execute(
                "SELECT COUNT(*) FROM legal_import_staging WHERE law=?",
                ("Synthetic staged law import",),
            ).fetchone()[0],
            1,
        )
        # Caller-supplied trust flags must not bypass the reviewed-corpus workflow.
        self.assertEqual(
            corpus.execute(
                "SELECT COUNT(*) FROM legal_documents WHERE official_title=?",
                ("Synthetic staged law import",),
            ).fetchone()[0],
            0,
        )

        control.execute(
            "DELETE FROM user_platform_roles WHERE user_id=? AND role_name='platform_reviewer'",
            (self.user_id,),
        )
        control.commit()
        revoked = self._mutate("post", "/law/import", json=payload, headers=headers)
        self.assertEqual(revoked.status_code, 403, revoked.text)
        self.assertEqual(
            corpus.execute(
                "SELECT COUNT(*) FROM legal_import_staging WHERE law=?",
                ("Synthetic staged law import",),
            ).fetchone()[0],
            1,
        )


    def test_assigned_tasks_do_not_bypass_case_membership_in_dashboard_or_calendar(self):
        owner_login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(owner_login.status_code, 200, owner_login.text)
        select_firm = self._mutate("post", "/session/firm", json={"firm_id": self.firm_a})
        self.assertEqual(select_firm.status_code, 200, select_firm.text)
        case_response = self._mutate(
            "post", "/cases", json={"title": "NONMEMBER_CASE_TITLE_SENTINEL"}
        )
        self.assertEqual(case_response.status_code, 200, case_response.text)
        case_id = case_response.json()["id"]

        nonmember_id = self.runtime.create_global_identity(
            "Case Nonmember", "case-nonmember@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(nonmember_id, self.firm_a, "lawyer")
        due = (date.today() + timedelta(days=10)).isoformat()
        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        tenant = self.runtime.firm_connection(context)
        self.assertEqual(
            tenant.execute("SELECT title FROM cases WHERE id=?", (case_id,)).fetchone()["title"],
            "NONMEMBER_CASE_TITLE_SENTINEL",
        )
        tenant.execute(
            "INSERT INTO tasks(case_id,title,assignee,due,status,created_by) VALUES(?,?,?,?,?,?)",
            (case_id, "NONMEMBER_ASSIGNED_TASK_SENTINEL", nonmember_id, due, "open", self.user_id),
        )
        tenant.commit()

        nonmember_client = TestClient(app.app)
        try:
            login = nonmember_client.post(
                "/login",
                json={"email": "case-nonmember@example.test", "password": self.password},
            )
            self.assertEqual(login.status_code, 200, login.text)
            dashboard = nonmember_client.get("/dashboard")
            self.assertEqual(dashboard.status_code, 200, dashboard.text)
            self.assertNotIn("NONMEMBER_ASSIGNED_TASK_SENTINEL", dashboard.text)
            self.assertNotIn("NONMEMBER_CASE_TITLE_SENTINEL", dashboard.text)

            calendar = nonmember_client.get("/calendar.ics")
            self.assertEqual(calendar.status_code, 200, calendar.text)
            self.assertEqual(calendar.headers.get("Cache-Control"), "no-store")
            self.assertEqual(calendar.headers.get("Pragma"), "no-cache")
            self.assertNotIn("NONMEMBER_ASSIGNED_TASK_SENTINEL", calendar.text)
            self.assertNotIn("NONMEMBER_CASE_TITLE_SENTINEL", calendar.text)
        finally:
            nonmember_client.close()

        # Firm admins retain access to the case and its task detail.
        admin_case = self.client.get(f"/cases/{case_id}")
        self.assertEqual(admin_case.status_code, 200, admin_case.text)
        self.assertIn("NONMEMBER_ASSIGNED_TASK_SENTINEL", admin_case.text)


    def test_application_responses_are_not_cacheable(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        responses = (
            self.client.get("/cases"),
            self.client.get("/audit/security-events"),
            self.client.get("/calendar.ics"),
            self.client.get("/cases/999999"),
        )
        self.assertEqual([response.status_code for response in responses], [200, 200, 200, 404])
        for response in responses:
            with self.subTest(path=response.request.url.path, status=response.status_code):
                self.assertEqual(response.headers.get("Cache-Control"), "no-store")
                self.assertEqual(response.headers.get("Pragma"), "no-cache")
                self.assertEqual(response.headers.get("Expires"), "0")


    def test_nonmember_cannot_infer_finance_deadline_or_hearing_data(self):
        owner_login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(owner_login.status_code, 200, owner_login.text)
        select_firm = self._mutate("post", "/session/firm", json={"firm_id": self.firm_a})
        self.assertEqual(select_firm.status_code, 200, select_firm.text)
        case_response = self._mutate(
            "post", "/cases", json={"title": "NONMEMBER_FINANCIAL_CASE_SENTINEL"}
        )
        self.assertEqual(case_response.status_code, 200, case_response.text)
        case_id = case_response.json()["id"]

        nonmember_id = self.runtime.create_global_identity(
            "Finance Nonmember", "finance-nonmember@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(nonmember_id, self.firm_a, "lawyer")
        due = (date.today() + timedelta(days=10)).isoformat()
        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        tenant = self.runtime.firm_connection(context)
        self.assertEqual(
            tenant.execute("SELECT title FROM cases WHERE id=?", (case_id,)).fetchone()["title"],
            "NONMEMBER_FINANCIAL_CASE_SENTINEL",
        )
        tenant.execute(
            "INSERT INTO fees(case_id,title,amount,paid,due,note,created_by) VALUES(?,?,?,?,?,?,?)",
            (case_id, "NONMEMBER_FEE_SENTINEL", 9876.50, 0, due, "", self.user_id),
        )
        tenant.execute(
            "INSERT INTO deadlines(case_id,title,due,note,created_by,status) VALUES(?,?,?,?,?,'open')",
            (case_id, "NONMEMBER_DEADLINE_SENTINEL", due, "", self.user_id),
        )
        tenant.execute(
            "INSERT INTO sessions(case_id,date,note,done) VALUES(?,?,?,0)",
            (case_id, due, "NONMEMBER_HEARING_SENTINEL",),
        )
        tenant.commit()

        nonmember_client = TestClient(app.app)
        try:
            login = nonmember_client.post(
                "/login",
                json={"email": "finance-nonmember@example.test", "password": self.password},
            )
            self.assertEqual(login.status_code, 200, login.text)

            dashboard = nonmember_client.get("/dashboard")
            self.assertEqual(dashboard.status_code, 200, dashboard.text)
            for sentinel in (
                "NONMEMBER_FINANCIAL_CASE_SENTINEL",
                "NONMEMBER_DEADLINE_SENTINEL",
                "NONMEMBER_HEARING_SENTINEL",
            ):
                self.assertNotIn(sentinel, dashboard.text)

            calendar = nonmember_client.get("/calendar.ics")
            self.assertEqual(calendar.status_code, 200, calendar.text)
            for sentinel in (
                "NONMEMBER_FINANCIAL_CASE_SENTINEL",
                "NONMEMBER_DEADLINE_SENTINEL",
                "NONMEMBER_HEARING_SENTINEL",
            ):
                self.assertNotIn(sentinel, calendar.text)

            finance = nonmember_client.get("/finance/overview")
            self.assertEqual(finance.status_code, 200, finance.text)
            self.assertEqual(finance.json()["fee_count"], 0)
            self.assertEqual(finance.json()["total_amount"], 0)
            self.assertEqual(finance.json()["outstanding_amount"], 0)
        finally:
            nonmember_client.close()


    def test_case_member_cannot_mutate_other_cases_tasks_or_deadlines_by_guessed_ids(self):
        owner_login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(owner_login.status_code, 200, owner_login.text)
        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_a})
        self.assertEqual(switched.status_code, 200, switched.text)
        case_a = self._mutate("post", "/cases", json={"title": "Authorized Case A"})
        case_b = self._mutate("post", "/cases", json={"title": "Private Case B"})
        self.assertEqual(case_a.status_code, 200, case_a.text)
        self.assertEqual(case_b.status_code, 200, case_b.text)
        case_a_id, case_b_id = case_a.json()["id"], case_b.json()["id"]

        tenant = self.runtime.firm_connection(
            self.runtime.resolve_context(self.user_id, self.firm_a)
        )
        nonmember_id = self.runtime.create_global_identity(
            "Case A Member", "case-a-member@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(nonmember_id, self.firm_a, "lawyer")
        tenant.execute(
            "INSERT INTO members(case_id,user_id,role) VALUES(?,?,?)",
            (case_a_id, nonmember_id, "lawyer"),
        )
        tenant.execute(
            "INSERT INTO tasks(case_id,title,assignee,status) VALUES(?,?,NULL,'open')",
            (case_b_id, "CASE_B_GUESSED_TASK_SENTINEL"),
        )
        task_id = tenant.execute(
            "SELECT id FROM tasks WHERE case_id=? AND title=?",
            (case_b_id, "CASE_B_GUESSED_TASK_SENTINEL"),
        ).fetchone()["id"]
        tenant.execute(
            "INSERT INTO deadlines(case_id,title,due,note,created_by,status) VALUES(?,?,?,?,?,'open')",
            (case_b_id, "CASE_B_GUESSED_DEADLINE_SENTINEL", "2026-11-10", "", self.user_id),
        )
        deadline_id = tenant.execute(
            "SELECT id FROM deadlines WHERE case_id=? AND title=?",
            (case_b_id, "CASE_B_GUESSED_DEADLINE_SENTINEL"),
        ).fetchone()["id"]
        tenant.commit()

        outsider = TestClient(app.app)
        try:
            login = outsider.post(
                "/login",
                json={"email": "case-a-member@example.test", "password": self.password},
            )
            self.assertEqual(login.status_code, 200, login.text)
            csrf = outsider.cookies.get("lc_csrf") or ""
            headers = {"X-CSRF-Token": csrf}
            switched = outsider.post(
                "/session/firm", json={"firm_id": self.firm_a}, headers=headers
            )
            self.assertEqual(switched.status_code, 200, switched.text)
            for path in (
                f"/tasks/{task_id}/claim",
                f"/tasks/{task_id}/done",
                f"/deadlines/{deadline_id}/done",
            ):
                with self.subTest(path=path):
                    response = outsider.post(path, headers=headers)
                    self.assertEqual(response.status_code, 404, response.text)
        finally:
            outsider.close()

        task = tenant.execute(
            "SELECT status,assignee FROM tasks WHERE id=?", (task_id,)
        ).fetchone()
        deadline = tenant.execute(
            "SELECT status FROM deadlines WHERE id=?", (deadline_id,)
        ).fetchone()
        self.assertEqual(task["status"], "open")
        self.assertIsNone(task["assignee"])
        self.assertEqual(deadline["status"], "open")

    def test_same_firm_nonmember_cannot_read_or_mutate_case_data(self):
        owner_login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(owner_login.status_code, 200, owner_login.text)
        select_firm = self._mutate("post", "/session/firm", json={"firm_id": self.firm_a})
        self.assertEqual(select_firm.status_code, 200, select_firm.text)
        case_response = self._mutate(
            "post", "/cases", json={"title": "SAME_FIRM_PRIVATE_CASE_SENTINEL"}
        )
        self.assertEqual(case_response.status_code, 200, case_response.text)
        case_id = case_response.json()["id"]
        document_response = self._mutate(
            "post",
            f"/cases/{case_id}/docs",
            json={"name": "private-evidence.txt", "text": "SAME_FIRM_PRIVATE_DOCUMENT_SENTINEL source text"},
        )
        self.assertEqual(document_response.status_code, 200, document_response.text)
        document_id = document_response.json()["document_id"]
        task_response = self._mutate(
            "post", f"/cases/{case_id}/tasks", json={"title": "PRIVATE_TASK_STATE_SENTINEL"}
        )
        self.assertEqual(task_response.status_code, 200, task_response.text)
        task_id = self.runtime.firm_connection(
            self.runtime.resolve_context(self.user_id, self.firm_a)
        ).execute(
            "SELECT id FROM tasks WHERE case_id=? AND title=?",
            (case_id, "PRIVATE_TASK_STATE_SENTINEL"),
        ).fetchone()["id"]
        owner_tenant = self.runtime.firm_connection(
            self.runtime.resolve_context(self.user_id, self.firm_a)
        )
        owner_tenant.execute(
            "INSERT INTO deadlines(case_id,title,due,note,created_by) VALUES(?,?,?,?,?)",
            (case_id, "PRIVATE_DEADLINE_STATE_SENTINEL", "2026-11-10", "", self.user_id),
        )
        owner_tenant.commit()
        deadline_id = owner_tenant.execute(
            "SELECT id FROM deadlines WHERE case_id=? AND title=?",
            (case_id, "PRIVATE_DEADLINE_STATE_SENTINEL"),
        ).fetchone()["id"]

        nonmember_id = self.runtime.create_global_identity(
            "Same Firm Nonmember", "same-firm-nonmember@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(nonmember_id, self.firm_a, "lawyer")
        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        tenant = self.runtime.firm_connection(context)
        baseline = {
            table: tenant.execute(f"SELECT COUNT(*) FROM {table} WHERE case_id=?", (case_id,)).fetchone()[0]
            for table in ("docs", "chunks", "notes", "deadlines", "fees", "sessions", "tasks", "events")
        }
        case_before = tenant.execute(
            "SELECT title,stage,status FROM cases WHERE id=?", (case_id,)
        ).fetchone()

        nonmember_client = TestClient(app.app)
        try:
            login = nonmember_client.post(
                "/login",
                json={"email": "same-firm-nonmember@example.test", "password": self.password},
            )
            self.assertEqual(login.status_code, 200, login.text)
            csrf = nonmember_client.cookies.get("lc_csrf") or ""
            headers = {"X-CSRF-Token": csrf}

            reads = (
                nonmember_client.get(f"/cases/{case_id}"),
                nonmember_client.get(f"/cases/{case_id}/docs/{document_id}/chunks"),
                nonmember_client.get(f"/cases/{case_id}/docs/{document_id}/reviews"),
            )
            for response in reads:
                with self.subTest(method="GET", path=response.request.url.path):
                    self.assertEqual(response.status_code, 404, response.text)
                    self.assertNotIn("SAME_FIRM_PRIVATE", response.text)

            mutations = (
                ("patch", f"/cases/{case_id}", {"stage": "closed"}),
                ("post", f"/cases/{case_id}/notes", {"text": "unauthorized note"}),
                ("post", f"/cases/{case_id}/deadlines", {"title": "unauthorized deadline", "due": "2026-11-10"}),
                ("post", f"/cases/{case_id}/fees", {"title": "unauthorized fee", "amount": 10, "paid": 0}),
                ("post", f"/cases/{case_id}/sessions", {"date": "2026-11-10", "note": "unauthorized session"}),
                ("post", f"/cases/{case_id}/tasks", {"title": "unauthorized task"}),
                ("post", f"/tasks/{task_id}/claim", {}),
                ("post", f"/tasks/{task_id}/done", {}),
                ("post", f"/deadlines/{deadline_id}/done", {}),
                ("post", f"/cases/{case_id}/docs", {"name": "unauthorized.txt", "text": "unauthorized document"}),
            )
            for method_name, path, payload in mutations:
                with self.subTest(method=method_name.upper(), path=path):
                    response = getattr(nonmember_client, method_name)(
                        path, json=payload, headers=headers
                    )
                    self.assertEqual(response.status_code, 404, response.text)

            upload = nonmember_client.post(
                f"/cases/{case_id}/upload",
                headers=headers,
                files={"file": ("unauthorized.txt", b"unauthorized upload text", "text/plain")},
            )
            self.assertEqual(upload.status_code, 404, upload.text)
            deletion = nonmember_client.delete(f"/cases/{case_id}", headers=headers)
            self.assertEqual(deletion.status_code, 404, deletion.text)
        finally:
            nonmember_client.close()

        task_after = tenant.execute(
            "SELECT status,assignee FROM tasks WHERE id=?", (task_id,)
        ).fetchone()
        self.assertEqual(task_after["status"], "open")
        self.assertIsNone(task_after["assignee"])
        deadline_after = tenant.execute(
            "SELECT status FROM deadlines WHERE id=?", (deadline_id,)
        ).fetchone()
        self.assertEqual(deadline_after["status"], "open")
        case_after = tenant.execute(
            "SELECT title,stage,status FROM cases WHERE id=?", (case_id,)
        ).fetchone()
        self.assertEqual(tuple(case_after), tuple(case_before))
        for table, count in baseline.items():
            with self.subTest(table=table):
                self.assertEqual(
                    tenant.execute(f"SELECT COUNT(*) FROM {table} WHERE case_id=?", (case_id,)).fetchone()[0],
                    count,
                )


    def test_document_persistence_rolls_back_metadata_chunks_and_event_together(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        select_firm = self._mutate("post", "/session/firm", json={"firm_id": self.firm_a})
        self.assertEqual(select_firm.status_code, 200, select_firm.text)
        case_response = self._mutate(
            "post", "/cases", json={"title": "Document Atomicity Case"}
        )
        self.assertEqual(case_response.status_code, 200, case_response.text)
        case_id = case_response.json()["id"]

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        tenant = self.runtime.firm_connection(context)
        before = {
            table: tenant.execute(f"SELECT COUNT(*) FROM {table} WHERE case_id=?", (case_id,)).fetchone()[0]
            for table in ("docs", "chunks", "events")
        }
        tenant.execute("""
            CREATE TRIGGER reject_document_audit_event
            BEFORE INSERT ON events
            WHEN NEW.text LIKE '%إضافة مستند:%'
            BEGIN SELECT RAISE(ABORT, 'document audit unavailable'); END
        """)
        tenant.commit()

        response = self._mutate(
            "post",
            f"/cases/{case_id}/docs",
            json={"name": "atomicity-test.txt", "text": "DOCUMENT_ATOMICITY_SENTINEL extracted source text"},
        )
        self.assertEqual(response.status_code, 503, response.text)
        for table, count in before.items():
            with self.subTest(table=table):
                self.assertEqual(
                    tenant.execute(f"SELECT COUNT(*) FROM {table} WHERE case_id=?", (case_id,)).fetchone()[0],
                    count,
                )
        self.assertFalse(tenant.in_transaction)


    def test_failed_matter_write_rolls_back_before_connection_is_reused(self):
        client = TestClient(app.app, raise_server_exceptions=False)
        try:
            login = client.post(
                "/login", json={"email": "OWNER@example.test", "password": self.password}
            )
            self.assertEqual(login.status_code, 200, login.text)
            csrf = client.cookies.get("lc_csrf") or ""
            headers = {"X-CSRF-Token": csrf}
            switched = client.post(
                "/session/firm", json={"firm_id": self.firm_a}, headers=headers
            )
            self.assertEqual(switched.status_code, 200, switched.text)
            case_response = client.post(
                "/cases", json={"title": "Failed Write Rollback Case"}, headers=headers
            )
            self.assertEqual(case_response.status_code, 200, case_response.text)
            case_id = case_response.json()["id"]

            context = self.runtime.resolve_context(self.user_id, self.firm_a)
            tenant = self.runtime.firm_connection(context)
            tenant.execute("""
                CREATE TRIGGER reject_note_event
                BEFORE INSERT ON events
                WHEN NEW.text LIKE '%إضافة ملاحظة%'
                BEGIN SELECT RAISE(ABORT, 'note audit unavailable'); END
            """)
            tenant.commit()

            failed = client.post(
                f"/cases/{case_id}/notes",
                json={"text": "FAILED_NOTE_MUST_NOT_PERSIST"},
                headers=headers,
            )
            self.assertEqual(failed.status_code, 503, failed.text)
            self.assertEqual(failed.headers.get("Cache-Control"), "no-store")

            tenant.execute("DROP TRIGGER reject_note_event")
            tenant.commit()
            succeeding = client.post(
                f"/cases/{case_id}/notes",
                json={"text": "successful follow-up note"},
                headers=headers,
            )
            self.assertEqual(succeeding.status_code, 200, succeeding.text)

            rows = tenant.execute(
                "SELECT text FROM notes WHERE case_id=? ORDER BY id", (case_id,)
            ).fetchall()
            self.assertEqual([row["text"] for row in rows], ["successful follow-up note"])
            events = tenant.execute(
                "SELECT text FROM events WHERE case_id=? ORDER BY id", (case_id,)
            ).fetchall()
            self.assertFalse(any("FAILED_NOTE_MUST_NOT_PERSIST" in row["text"] for row in events))
        finally:
            client.close()


    def test_failed_shared_corpus_import_rolls_back_partial_batch(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        control = self.runtime.control_connection()
        control.execute(
            "INSERT INTO user_platform_roles(user_id,role_name) VALUES(?,'platform_reviewer')",
            (self.user_id,),
        )
        control.commit()

        corpus = self.runtime.legal_corpus_connection()
        corpus.execute("""
            CREATE TRIGGER reject_second_staged_import
            BEFORE INSERT ON legal_import_staging
            WHEN NEW.law='Atomic staged import' AND NEW.article='2'
            BEGIN SELECT RAISE(ABORT, 'staging insert unavailable'); END
        """)
        corpus.commit()
        payload = [
            {"law": "Atomic staged import", "article": "1", "text": "first batch row", "source": "TEST"},
            {"law": "Atomic staged import", "article": "2", "text": "second batch row", "source": "TEST"},
        ]
        headers = {"X-Admin-Token": self.admin_token}
        failed = self._mutate("post", "/law/import", json=payload, headers=headers)
        self.assertEqual(failed.status_code, 503, failed.text)
        self.assertEqual(
            corpus.execute(
                "SELECT COUNT(*) FROM legal_import_staging WHERE law=?",
                ("Atomic staged import",),
            ).fetchone()[0],
            0,
        )

        corpus.execute("DROP TRIGGER reject_second_staged_import")
        corpus.commit()
        succeeding = self._mutate(
            "post",
            "/law/import",
            json=[{"law": "Atomic staged import", "article": "3", "text": "successful follow-up row", "source": "TEST"}],
            headers=headers,
        )
        self.assertEqual(succeeding.status_code, 200, succeeding.text)
        rows = corpus.execute(
            "SELECT article FROM legal_import_staging WHERE law=? ORDER BY article",
            ("Atomic staged import",),
        ).fetchall()
        self.assertEqual([row["article"] for row in rows], ["3"])


    def test_tenant_and_corpus_mutation_routes_declare_rollback_guards(self):
        mutation_methods = {"POST", "PUT", "PATCH", "DELETE"}
        tenant_prefixes = ("/cases", "/tasks", "/deadlines", "/research/conversations")
        corpus_write_paths = {
            "/law/import",
            "/legal-review/versions/{version_id}/events",
        }
        seen_tenant = set()
        seen_corpus = set()
        for route in app.app.routes:
            path = getattr(route, "path", "")
            methods = set(getattr(route, "methods", set())) & mutation_methods
            if not methods:
                continue
            domain = getattr(route.endpoint, "_rollback_transaction_domain", None)
            if path.startswith(tenant_prefixes):
                seen_tenant.update((method, path) for method in methods)
                with self.subTest(methods=sorted(methods), path=path):
                    self.assertEqual(domain, "tenant")
            if path in corpus_write_paths:
                seen_corpus.update((method, path) for method in methods)
                with self.subTest(methods=sorted(methods), path=path):
                    self.assertEqual(domain, "corpus")

        self.assertTrue(seen_tenant, "expected to inspect tenant mutation routes")
        self.assertIn(("POST", "/cases"), seen_tenant)
        self.assertIn(("POST", "/research/conversations"), seen_tenant)
        self.assertEqual(seen_corpus, {
            ("POST", "/law/import"),
            ("POST", "/legal-review/versions/{version_id}/events"),
        })


    def test_same_firm_research_conversations_are_private_to_their_owner(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/research/conversations")
        self.assertEqual(created.status_code, 200, created.text)
        conversation_id = created.json()["id"]
        renamed = self._mutate(
            "patch",
            f"/research/conversations/{conversation_id}",
            json={"title": "SAME_FIRM_PRIVATE_RESEARCH_SENTINEL"},
        )
        self.assertEqual(renamed.status_code, 200, renamed.text)

        nonmember_id = self.runtime.create_global_identity(
            "Research Colleague", "research-colleague@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(nonmember_id, self.firm_a, "lawyer")
        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        tenant = self.runtime.firm_connection(context)

        nonmember_client = TestClient(app.app)
        try:
            nonmember_login = nonmember_client.post(
                "/login",
                json={"email": "research-colleague@example.test", "password": self.password},
            )
            self.assertEqual(nonmember_login.status_code, 200, nonmember_login.text)
            csrf = nonmember_client.cookies.get("lc_csrf") or ""
            headers = {"X-CSRF-Token": csrf}

            listing = nonmember_client.get("/research/conversations")
            self.assertEqual(listing.status_code, 200, listing.text)
            self.assertNotIn("SAME_FIRM_PRIVATE_RESEARCH_SENTINEL", listing.text)
            detail = nonmember_client.get(f"/research/conversations/{conversation_id}")
            self.assertEqual(detail.status_code, 404, detail.text)
            rename_attempt = nonmember_client.patch(
                f"/research/conversations/{conversation_id}",
                json={"title": "UNAUTHORIZED_RESEARCH_RENAME"},
                headers=headers,
            )
            self.assertEqual(rename_attempt.status_code, 404, rename_attempt.text)
            message_attempt = nonmember_client.post(
                f"/research/conversations/{conversation_id}/messages",
                json={"message": "UNAUTHORIZED_RESEARCH_MESSAGE", "scope": "verified_corpus"},
                headers=headers,
            )
            self.assertEqual(message_attempt.status_code, 404, message_attempt.text)
            save_attempt = nonmember_client.post(
                f"/research/conversations/{conversation_id}/save-to-matter",
                json={"case_id": 1, "message_id": 1},
                headers=headers,
            )
            self.assertEqual(save_attempt.status_code, 404, save_attempt.text)
            delete_attempt = nonmember_client.delete(
                f"/research/conversations/{conversation_id}", headers=headers
            )
            self.assertEqual(delete_attempt.status_code, 404, delete_attempt.text)
        finally:
            nonmember_client.close()

        row = tenant.execute(
            "SELECT title,user_id,firm_id FROM research_conversations WHERE id=?",
            (conversation_id,),
        ).fetchone()
        self.assertEqual(row["title"], "SAME_FIRM_PRIVATE_RESEARCH_SENTINEL")
        self.assertEqual(row["user_id"], self.user_id)
        self.assertEqual(row["firm_id"], self.firm_a)
        self.assertEqual(
            tenant.execute(
                "SELECT COUNT(*) FROM research_messages WHERE conversation_id=?",
                (conversation_id,),
            ).fetchone()[0],
            0,
        )

        owner_rename = self._mutate(
            "patch", f"/research/conversations/{conversation_id}",
            json={"title": "  Owner research title  "},
        )
        self.assertEqual(owner_rename.status_code, 200, owner_rename.text)
        self.assertEqual(owner_rename.json()["title"], "Owner research title")
        blank_rename = self._mutate(
            "patch", f"/research/conversations/{conversation_id}",
            json={"title": "   "},
        )
        self.assertEqual(blank_rename.status_code, 422, blank_rename.text)
        deleted = self._mutate("delete", f"/research/conversations/{conversation_id}")
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertEqual(
            self.client.get(f"/research/conversations/{conversation_id}").status_code,
            404,
        )
        self.assertFalse(any(
            item["id"] == conversation_id
            for item in self.client.get("/research/conversations").json()
        ))


    def test_failed_legal_review_event_write_rolls_back_and_returns_503(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        corpus = self.runtime.legal_corpus_connection()
        corpus.execute(
            "INSERT INTO legal_sources(source_id,authority,source_page) VALUES(?,?,?)",
            ("review-rollback-source", "Review rollback test", "https://example.test/source"),
        )
        corpus.execute(
            "INSERT INTO legal_documents(document_id,source_id,official_title,document_type) VALUES(?,?,?,?)",
            ("review-rollback-doc", "review-rollback-source", "Review rollback document", "statute"),
        )
        corpus.execute(
            "INSERT INTO legal_document_versions(version_id,document_id,version_label,source_url) VALUES(?,?,?,?)",
            (780, "review-rollback-doc", "v1", "https://example.test/source/v1"),
        )
        corpus.commit()

        control = self.runtime.control_connection()
        control.execute(
            "INSERT INTO user_platform_roles(user_id,role_name) VALUES(?,'platform_reviewer')",
            (self.user_id,),
        )
        control.commit()
        corpus.execute("""
            CREATE TRIGGER reject_review_event_780
            BEFORE INSERT ON legal_review_events
            WHEN NEW.version_id=780
            BEGIN SELECT RAISE(ABORT, 'review event unavailable'); END
        """)
        corpus.commit()

        endpoint = "/legal-review/versions/780/events"
        payload = {
            "review_type": "legal_status",
            "outcome": "pending",
            "notes": "Synthetic rollback regression",
        }
        failed = self._mutate("post", endpoint, json=payload)
        self.assertEqual(failed.status_code, 503, failed.text)
        self.assertEqual(
            corpus.execute(
                "SELECT COUNT(*) FROM legal_review_events WHERE version_id=780"
            ).fetchone()[0],
            0,
        )

        corpus.execute("DROP TRIGGER reject_review_event_780")
        corpus.commit()
        succeeding = self._mutate("post", endpoint, json=payload)
        self.assertEqual(succeeding.status_code, 201, succeeding.text)
        self.assertEqual(
            corpus.execute(
                "SELECT COUNT(*) FROM legal_review_events WHERE version_id=780"
            ).fetchone()[0],
            1,
        )


    def test_unhandled_errors_return_generic_non_cacheable_responses(self):
        original_rate_limit = app._rate_limit
        client = TestClient(app.app, raise_server_exceptions=False)
        def fail_rate_limit(_request):
            raise RuntimeError("SENSITIVE_ERROR_SENTINEL")
        app._rate_limit = fail_rate_limit
        try:
            response = client.get("/")
            self.assertEqual(response.status_code, 500)
            self.assertEqual(response.json(), {"detail": "internal server error"})
            self.assertNotIn("SENSITIVE_ERROR_SENTINEL", response.text)
            self.assertEqual(response.headers.get("Cache-Control"), "no-store")
            self.assertEqual(response.headers.get("Pragma"), "no-cache")
            self.assertEqual(response.headers.get("Expires"), "0")
        finally:
            app._rate_limit = original_rate_limit
            client.close()


    def test_case_member_removal_requires_case_lead_and_clears_task_assignment(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate(
            "post", "/cases", json={"title": "Membership removal authorization matter"}
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        member_id = self.runtime.create_global_identity(
            "Removal Target", "removal-target@example.test", app.hash_pw(self.password)
        )
        self.runtime.create_firm_membership(member_id, self.firm_a, "lawyer")
        added = self._mutate(
            "post", f"/cases/{case_id}/members",
            json={"user_id": member_id, "role": "associate"},
        )
        self.assertEqual(added.status_code, 200, added.text)
        task = self._mutate(
            "post", f"/cases/{case_id}/tasks",
            json={"title": "Assigned task must be cleared", "assignee": member_id},
        )
        self.assertEqual(task.status_code, 200, task.text)

        context_a = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_a_db = self.runtime.firm_connection(context_a)
        task_id = firm_a_db.execute(
            "SELECT id FROM tasks WHERE case_id=? AND assignee=?",
            (case_id, member_id),
        ).fetchone()["id"]

        # A same-firm case member is not a case lead and cannot remove another member.
        member_client = TestClient(app.app)
        try:
            member_login = member_client.post(
                "/login",
                json={"email": "removal-target@example.test", "password": self.password},
            )
            self.assertEqual(member_login.status_code, 200, member_login.text)
            csrf = member_client.cookies.get("lc_csrf")
            denied = member_client.delete(
                f"/cases/{case_id}/members/{self.user_id}",
                headers={"X-CSRF-Token": csrf or ""},
            )
            self.assertEqual(denied.status_code, 404, denied.text)
        finally:
            member_client.close()

        # A firm switch must not resolve or mutate the original firm's membership.
        switched = self._mutate("post", "/session/firm", json={"firm_id": self.firm_b})
        self.assertEqual(switched.status_code, 200, switched.text)
        cross_firm = self._mutate("delete", f"/cases/{case_id}/members/{member_id}")
        self.assertEqual(cross_firm.status_code, 404, cross_firm.text)
        self.assertEqual(
            firm_a_db.execute(
                "SELECT COUNT(*) FROM members WHERE case_id=? AND user_id=?",
                (case_id, member_id),
            ).fetchone()[0],
            1,
        )
        self.assertEqual(
            firm_a_db.execute("SELECT assignee FROM tasks WHERE id=?", (task_id,)).fetchone()[0],
            member_id,
        )

        switched_back = self._mutate("post", "/session/firm", json={"firm_id": self.firm_a})
        self.assertEqual(switched_back.status_code, 200, switched_back.text)

        # If the audit event cannot be written, removal and task unassignment must
        # roll back together rather than leaving a partially applied operation.
        with patch.object(app, "log", side_effect=RuntimeError("injected audit failure")):
            failed_removal = self._mutate("delete", f"/cases/{case_id}/members/{member_id}")
        self.assertEqual(failed_removal.status_code, 500, failed_removal.text)
        # End any read snapshot on the observer connection before checking the
        # committed state written by the request connection.
        firm_a_db.commit()
        self.assertEqual(
            firm_a_db.execute(
                "SELECT COUNT(*) FROM members WHERE case_id=? AND user_id=?",
                (case_id, member_id),
            ).fetchone()[0],
            1,
        )
        self.assertEqual(
            firm_a_db.execute("SELECT assignee FROM tasks WHERE id=?", (task_id,)).fetchone()[0],
            member_id,
        )

        removed = self._mutate("delete", f"/cases/{case_id}/members/{member_id}")
        self.assertEqual(removed.status_code, 200, removed.text)
        firm_a_db.commit()
        self.assertEqual(
            firm_a_db.execute(
                "SELECT COUNT(*) FROM members WHERE case_id=? AND user_id=?",
                (case_id, member_id),
            ).fetchone()[0],
            0,
        )
        self.assertIsNone(
            firm_a_db.execute("SELECT assignee FROM tasks WHERE id=?", (task_id,)).fetchone()[0],
            "removing a case member must clear their task assignments atomically",
        )


    def test_firm_document_search_scope_requires_case_and_preserves_page_pinpoint(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate(
            "post", "/cases", json={"title": "Firm-document search scope test"}
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]

        uploaded = self._mutate(
            "post", f"/cases/{case_id}/docs",
            json={
                "name": "page-pinpoint-fixture.txt",
                "text": "[[PDF_PAGE:17]]\n\nThe limitation period in this example is seven years.",
            },
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        chunk = firm_db.execute(
            "SELECT id FROM chunks WHERE document_id=? AND case_id=? ORDER BY id LIMIT 1",
            (uploaded.json()["document_id"], case_id),
        ).fetchone()
        self.assertIsNotNone(chunk)
        firm_db.execute(
            "UPDATE chunks SET ref=?,article=? WHERE id=?",
            ("Page pinpoint fixture", "12", chunk["id"]),
        )
        firm_db.commit()

        conversation = self._mutate("post", "/research/conversations", json={})
        self.assertEqual(conversation.status_code, 200, conversation.text)
        conversation_id = conversation.json()["id"]
        missing_case = self._mutate(
            "post", f"/research/conversations/{conversation_id}/messages",
            json={"message": "limitation period example", "scope": "firm_documents"},
        )
        self.assertEqual(missing_case.status_code, 422, missing_case.text)

        firm_scope = self._mutate(
            "post", f"/research/conversations/{conversation_id}/messages",
            json={
                "message": "مدة التقادم",
                "scope": "firm_documents",
                "case_id": case_id,
            },
        )
        self.assertEqual(firm_scope.status_code, 200, firm_scope.text)
        self.assertEqual(firm_scope.json()["scope"], "firm_documents")
        self.assertTrue(firm_scope.json().get("sources"), firm_scope.json())
        self.assertTrue(all(
            source.get("case_id") == case_id
            for source in firm_scope.json().get("sources", [])
        ))
        self.assertTrue(all(
            source.get("source_status") == "user_supplied_unverified"
            for source in firm_scope.json().get("sources", [])
        ))

        page_search = self._mutate(
            "post", "/search",
            json={
                "q": "limitation period example",
                "case_id": case_id,
                "scope": "firm_documents",
            },
        )
        self.assertEqual(page_search.status_code, 200, page_search.text)
        page_hit = next(
            item for item in page_search.json()
            if item["ref"] == "Page pinpoint fixture"
        )
        self.assertEqual(page_hit["page"], "17")
        self.assertEqual(page_hit["article"], "12")
        self.assertIn("limitation", page_hit["matched_terms"])
        self.assertIn("Page 17", page_hit["citation"])
        self.assertIn("Article 12", page_hit["citation"])

        verified_scope = self._mutate(
            "post", "/search",
            json={"q": "التقادم المثال التشغيلي", "scope": "verified_corpus"},
        )
        self.assertEqual(verified_scope.status_code, 200, verified_scope.text)
        self.assertFalse(any(
            item.get("ref") == "Page pinpoint fixture"
            for item in verified_scope.json()
        ))


    def test_research_discovery_escapes_like_wildcards_and_returns_metadata_only(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        corpus = self.runtime.legal_corpus_connection()

        no_match = self._mutate(
            "post", "/research/discover",
            json={"q": "zz_no_catalog_match_2026"},
        )
        self.assertEqual(no_match.status_code, 200, no_match.text)
        self.assertEqual(no_match.json()["scope"], "metadata_only")
        self.assertFalse(no_match.json()["found"])
        self.assertEqual(no_match.json()["results"], [])
        self.assertIn("does not establish", no_match.json()["message"])
        self.assertIn("No source page was scraped or downloaded", no_match.json()["warning"])
        whitespace = self._mutate(
            "post", "/research/discover", json={"q": "   "}
        )
        self.assertEqual(whitespace.status_code, 422, whitespace.text)

        corpus.execute(
            "INSERT INTO legal_sources(source_id,authority,source_page) VALUES(?,?,?)",
            ("literal-query-source", "Literal query test authority", "https://example.test/metadata"),
        )
        corpus.executemany(
            "INSERT INTO legal_documents(document_id,source_id,official_title,document_type) VALUES(?,?,?,?)",
            [
                (
                    "literal-query-target",
                    "literal-query-source",
                    "Case_A_100%_Contract",
                    "synthetic-metadata",
                ),
                (
                    "literal-query-wildcard-decoy",
                    "literal-query-source",
                    "CaseXA1000XContract",
                    "synthetic-metadata",
                ),
            ],
        )
        corpus.commit()

        response = self._mutate(
            "post", "/research/discover",
            json={"q": "Case_A_100%_Contract", "as_of": "2026-10-28"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["scope"], "metadata_only")
        self.assertEqual(
            [item["document_id"] for item in body["results"]],
            ["literal-query-target"],
            "underscore and percent characters must be treated as literal search input, not LIKE wildcards",
        )
        result = body["results"][0]
        self.assertFalse(result["text_available"])
        self.assertEqual(result["status"], "metadata_only")
        self.assertIn("no current-law status is implied", result["warning"])
        self.assertNotIn("text", result)
        self.assertNotIn("content", result)
        self.assertNotIn("passage", result)
        self.assertNotIn("CaseXA1000XContract", response.text)


    def test_research_save_to_matter_requires_case_membership_and_rolls_back_audit_failure(self):
        owner_login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(owner_login.status_code, 200, owner_login.text)
        created_case = self._mutate(
            "post", "/cases", json={"title": "Research memo membership boundary matter"}
        )
        self.assertEqual(created_case.status_code, 200, created_case.text)
        case_id = created_case.json()["id"]

        researcher_id = self.runtime.create_global_identity(
            "Researcher Without Matter Access",
            "researcher-without-matter-access@example.test",
            app.hash_pw(self.password),
        )
        self.runtime.create_firm_membership(researcher_id, self.firm_a, "lawyer")
        researcher = TestClient(app.app)
        try:
            login = researcher.post(
                "/login",
                json={"email": "researcher-without-matter-access@example.test", "password": self.password},
            )
            self.assertEqual(login.status_code, 200, login.text)
            csrf = researcher.cookies.get("lc_csrf") or ""
            headers = {"X-CSRF-Token": csrf}

            created_conversation = researcher.post(
                "/research/conversations", headers=headers
            )
            self.assertEqual(created_conversation.status_code, 200, created_conversation.text)
            conversation_id = created_conversation.json()["id"]
            context = self.runtime.resolve_context(researcher_id, self.firm_a)
            tenant = self.runtime.firm_connection(context)
            message_json = json.dumps({
                "answer": "RESEARCH_MEMO_PRIVATE_ANSWER_SENTINEL",
                "requested_as_of": "2026-10-28",
                "scope": "verified_corpus",
                "sources": [],
            })
            message_id = tenant.execute(
                "INSERT INTO research_messages(conversation_id,role,content,message_json) "
                "VALUES(?,?,?,?)",
                (conversation_id, "assistant", "RESEARCH_MEMO_PRIVATE_ANSWER_SENTINEL", message_json),
            ).lastrowid
            tenant.commit()

            notes_before = tenant.execute(
                "SELECT COUNT(*) FROM notes WHERE case_id=?", (case_id,)
            ).fetchone()[0]
            events_before = tenant.execute(
                "SELECT COUNT(*) FROM events WHERE case_id=?", (case_id,)
            ).fetchone()[0]
            denied = researcher.post(
                f"/research/conversations/{conversation_id}/save-to-matter",
                json={"case_id": case_id, "message_id": message_id},
                headers=headers,
            )
            self.assertEqual(denied.status_code, 404, denied.text)
            tenant.commit()
            self.assertEqual(
                tenant.execute("SELECT COUNT(*) FROM notes WHERE case_id=?", (case_id,)).fetchone()[0],
                notes_before,
            )
            self.assertEqual(
                tenant.execute("SELECT COUNT(*) FROM events WHERE case_id=?", (case_id,)).fetchone()[0],
                events_before,
            )

            # Grant explicit case membership, then prove the note and audit event
            # remain atomic if the event write fails.
            added = self._mutate(
                "post", f"/cases/{case_id}/members",
                json={"user_id": researcher_id, "role": "associate"},
            )
            self.assertEqual(added.status_code, 200, added.text)
            # Membership creation itself writes an audit event, so reset the
            # baseline after that authorized mutation before testing rollback.
            tenant.commit()
            notes_before = tenant.execute(
                "SELECT COUNT(*) FROM notes WHERE case_id=?", (case_id,)
            ).fetchone()[0]
            events_before = tenant.execute(
                "SELECT COUNT(*) FROM events WHERE case_id=?", (case_id,)
            ).fetchone()[0]
            with patch.object(app, "log", side_effect=RuntimeError("injected research memo audit failure")):
                failed = researcher.post(
                    f"/research/conversations/{conversation_id}/save-to-matter",
                    json={"case_id": case_id, "message_id": message_id},
                    headers=headers,
                )
            self.assertEqual(failed.status_code, 500, failed.text)
            tenant.commit()
            self.assertEqual(
                tenant.execute("SELECT COUNT(*) FROM notes WHERE case_id=?", (case_id,)).fetchone()[0],
                notes_before,
            )
            self.assertEqual(
                tenant.execute("SELECT COUNT(*) FROM events WHERE case_id=?", (case_id,)).fetchone()[0],
                events_before,
            )

            saved = researcher.post(
                f"/research/conversations/{conversation_id}/save-to-matter",
                json={"case_id": case_id, "message_id": message_id},
                headers=headers,
            )
            self.assertEqual(saved.status_code, 201, saved.text)
            tenant.commit()
            self.assertEqual(
                tenant.execute("SELECT COUNT(*) FROM notes WHERE case_id=?", (case_id,)).fetchone()[0],
                notes_before + 1,
            )
            self.assertEqual(
                tenant.execute("SELECT COUNT(*) FROM events WHERE case_id=?", (case_id,)).fetchone()[0],
                events_before + 1,
            )
            note = tenant.execute(
                "SELECT text FROM notes WHERE id=?", (saved.json()["note_id"],)
            ).fetchone()
            self.assertIn("RESEARCH_MEMO_PRIVATE_ANSWER_SENTINEL", note["text"])
            self.assertIn("LEGAL RESEARCH MEMO", note["text"])
            self.assertIn("As of: 2026-10-28", note["text"])
            self.assertIn("Search scope: verified_corpus", note["text"])
            self.assertIn("not evidence that no applicable law exists", note["text"])
        finally:
            researcher.close()


    def test_security_headers_and_csp_nonce_are_present(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers.get("X-Content-Type-Options"), "nosniff")
        self.assertEqual(response.headers.get("X-Frame-Options"), "DENY")
        policy = response.headers.get("Content-Security-Policy", "")
        self.assertIn("script-src", policy)
        script_policy = policy.split("script-src", 1)[1].split(";", 1)[0]
        self.assertNotIn("unsafe-inline", script_policy)
        self.assertIn("nonce-", script_policy)
        self.assertIn("<script nonce=", response.text)


    def test_firm_workflow_claim_calendar_and_fee_aggregates(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate(
            "post", "/cases", json={"title": "Workflow acceptance matter", "client": "Workflow client"}
        )
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        due = (date.today() + timedelta(days=7)).isoformat()

        task = self._mutate(
            "post", f"/cases/{case_id}/tasks",
            json={"title": "Claimable workflow task", "due": due},
        )
        self.assertEqual(task.status_code, 200, task.text)
        db = self.runtime.firm_connection(
            self.runtime.resolve_context(self.user_id, self.firm_a)
        )
        task_id = db.execute(
            "SELECT id FROM tasks WHERE case_id=? AND title=?",
            (case_id, "Claimable workflow task"),
        ).fetchone()["id"]
        dashboard = self.client.get("/dashboard")
        self.assertEqual(dashboard.status_code, 200, dashboard.text)
        self.assertTrue(any(
            item["title"] == "Claimable workflow task"
            for item in dashboard.json()["unassigned_tasks"]
        ))
        claimed = self._mutate("post", f"/tasks/{task_id}/claim")
        self.assertEqual(claimed.status_code, 200, claimed.text)
        duplicate_claim = self._mutate("post", f"/tasks/{task_id}/claim")
        self.assertEqual(duplicate_claim.status_code, 409, duplicate_claim.text)
        dashboard_after_claim = self.client.get("/dashboard")
        self.assertEqual(dashboard_after_claim.status_code, 200, dashboard_after_claim.text)
        self.assertTrue(any(
            item["title"] == "Claimable workflow task"
            for item in dashboard_after_claim.json()["tasks"]
        ))
        self.assertFalse(any(
            item["title"] == "Claimable workflow task"
            for item in dashboard_after_claim.json()["unassigned_tasks"]
        ))

        deadline = self._mutate(
            "post", f"/cases/{case_id}/deadlines",
            json={"title": "Workflow appeal deadline", "due": due, "note": "Verify source"},
        )
        self.assertEqual(deadline.status_code, 200, deadline.text)
        hearing = self._mutate(
            "post", f"/cases/{case_id}/sessions",
            json={"date": due, "note": "Workflow hearing"},
        )
        self.assertEqual(hearing.status_code, 200, hearing.text)

        fee = self._mutate(
            "post", f"/cases/{case_id}/fees",
            json={"title": "Workflow retainer", "amount": 1000, "paid": 250, "due": due},
        )
        self.assertEqual(fee.status_code, 200, fee.text)
        finance = self.client.get("/finance/overview")
        self.assertEqual(finance.status_code, 200, finance.text)
        self.assertEqual(finance.json()["fee_count"], 1)
        self.assertEqual(finance.json()["total_amount"], 1000)
        self.assertEqual(finance.json()["total_paid"], 250)
        self.assertEqual(finance.json()["outstanding_amount"], 750)
        self.assertEqual(finance.json()["overdue_count"], 0)

        invalid_fee = self._mutate(
            "post", f"/cases/{case_id}/fees",
            json={"title": "Invalid fee", "amount": 100, "paid": 101},
        )
        self.assertEqual(invalid_fee.status_code, 422, invalid_fee.text)
        negative_fee = self._mutate(
            "post", f"/cases/{case_id}/fees",
            json={"title": "Negative fee", "amount": -1, "paid": 0},
        )
        self.assertEqual(negative_fee.status_code, 422, negative_fee.text)

        calendar = self.client.get("/calendar.ics")
        self.assertEqual(calendar.status_code, 200, calendar.text)
        self.assertIn("Task: Claimable workflow task", calendar.text)
        self.assertIn("Legal deadline: Workflow appeal deadline", calendar.text)
        self.assertIn("Hearing: Workflow acceptance matter", calendar.text)


    def test_calendar_text_escaping_prevents_crlf_line_injection(self):
        self.assertEqual(
            app._ics_escape("a\\b,c;d\r\ne\rf\ng"),
            "a\\\\b\\,c\\;d\\ne\\nf\\ng",
        )


    def test_candidate_catalogue_dashboard_denominator_requires_unique_stable_ids(self):
        catalogue = {
            "documents": [
                {"document_id": "LAW-001", "practice_areas": ["Labour", "employment law"]},
                {"document_id": " LAW-001 ", "practice_area": "civil law"},
                {"document_id": "LAW-002", "practice_area": "criminal law"},
                {"document_id": "  ", "practice_area": "tax law"},
                {"practice_area": "family law"},
                {"document_id": "LAW-003", "practice_areas": "not-a-list"},
                None,
            ]
        }
        path = self.root / "candidate-catalogue.json"
        path.write_text(json.dumps(catalogue), encoding="utf-8")
        with patch.object(app, "LEGAL_CANDIDATE_CATALOG_PATH", path):
            result = app._candidate_catalogue_practice_area_counts()

        self.assertTrue(result["available"])
        self.assertEqual(result["candidate_instrument_count"], 3)
        self.assertEqual(result["practice_area_counts"]["labour"], 1)
        self.assertEqual(result["practice_area_counts"]["criminal"], 1)
        self.assertEqual(result["practice_area_counts"]["unknown"], 1)
        self.assertEqual(result["invalid_entries"], 5)

    def test_legal_corpus_coverage_counts_only_permission_complete_versions(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        control = self.runtime.control_connection()
        control.execute(
            "INSERT INTO user_platform_roles(user_id,role_name) VALUES(?,?)",
            (self.user_id, "platform_reviewer"),
        )
        control.commit()

        before = self.client.get("/legal-review/coverage")
        self.assertEqual(before.status_code, 200, before.text)
        verified_before = before.json()["metrics"]["verified_operational_versions"]

        corpus = self.runtime.legal_corpus_connection()
        corpus.execute(
            "INSERT INTO legal_sources(source_id,authority,source_page,access_status,"
            "collection_permission,commercial_reuse_permission,permission_evidence) "
            "VALUES(?,?,?,?,?,?,?)",
            ("RUNTIME-COVERAGE-SOURCE", "Coverage test authority",
             "https://example.test/coverage-source", "reviewed", "allowed", "allowed",
             "synthetic source permission evidence"),
        )
        corpus.execute(
            "INSERT INTO legal_documents(document_id,source_id,official_title,document_type,practice_area) "
            "VALUES(?,?,?,?,?)",
            ("RUNTIME-COVERAGE-DOC", "RUNTIME-COVERAGE-SOURCE",
             "Coverage test instrument", "law", "labour"),
        )
        version_cursor = corpus.execute(
            "INSERT INTO legal_document_versions(document_id,version_label,source_url,publication_date,"
            "effective_date,legal_status,review_status,authenticity_verified,status_verified,"
            "amendment_history_verified,permission_evidence,original_sha256) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
            ("RUNTIME-COVERAGE-DOC", "verified-version", "https://example.test/coverage-v1",
             "2020-01-01", "2020-01-01", "in_force", "verified", 1, 1, 1,
             "synthetic version permission evidence", "a" * 64),
        )
        version_id = version_cursor.lastrowid
        for review_type in ("source_authenticity", "extraction_quality", "legal_status", "amendment_history", "licensing", "lawyer_approval"):
            notes = "Synthetic review completed."
            if review_type == "lawyer_approval":
                notes = "[qualified_egyptian_counsel_attested:user_id=42] " + notes
            corpus.execute(
                "INSERT INTO legal_review_events(version_id,reviewer,review_type,outcome,notes,evidence_url) VALUES(?,?,?,?,?,?)",
                (version_id, "Synthetic reviewer", review_type, "pass", notes, "https://example.test/review"),
            )
        corpus.execute(
            "INSERT INTO chunks(ref,article,page,text,norm,document_version_id,in_force,as_of,source,verified) "
            "VALUES(?,?,?,?,?,?,?,?,?,?)",
            ("Coverage test instrument - article 1", "1", None, "Coverage legal text",
             app.norm("Coverage legal text"), version_id, 1, "2020-01-01",
             "https://example.test/coverage-source", 1),
        )
        corpus.execute(
            "INSERT INTO legal_document_versions(document_id,version_label,source_url,publication_date,"
            "effective_date,legal_status,review_status,authenticity_verified,status_verified,"
            "amendment_history_verified) VALUES(?,?,?,?,?,?,?,?,?,?)",
            ("RUNTIME-COVERAGE-DOC", "missing-permission-evidence",
             "https://example.test/coverage-v2", "2021-01-01", "2021-01-01",
             "in_force", "verified", 1, 1, 1),
        )
        corpus.execute(
            "INSERT INTO legal_sources(source_id,authority,source_page) VALUES(?,?,?)",
            ("RUNTIME-COVERAGE-METADATA-SOURCE", "Metadata-only authority",
             "https://example.test/metadata-only"),
        )
        corpus.execute(
            "INSERT INTO legal_documents(document_id,source_id,official_title,document_type) "
            "VALUES(?,?,?,?)",
            ("RUNTIME-COVERAGE-METADATA-DOC", "RUNTIME-COVERAGE-METADATA-SOURCE",
             "Metadata-only instrument", "law"),
        )
        corpus.commit()

        rights_evidence = {
            "source_permission_register_match": True,
            "source_permission_register_collection_permission": "allowed",
            "source_permission_register_commercial_reuse_permission": "allowed",
            "permission_evidence_url": "https://example.test/rights/collection",
            "commercial_reuse_evidence_url": "https://example.test/rights/commercial",
            "terms_checked_on": date.today().isoformat(),
            "commercial_terms_checked_on": date.today().isoformat(),
        }
        def permission_record(source_id, source_page):
            if source_id == "RUNTIME-COVERAGE-SOURCE":
                return dict(rights_evidence)
            return {
                "source_permission_register_match": False,
                "source_permission_register_collection_permission": "unknown",
                "source_permission_register_commercial_reuse_permission": "unknown",
                "permission_evidence_url": None,
                "commercial_reuse_evidence_url": None,
                "terms_checked_on": None,
                "commercial_terms_checked_on": None,
            }
        with patch.object(app, "_central_source_permission_evidence", side_effect=permission_record):
            response = self.client.get("/legal-review/coverage")
            queue = self.client.get("/legal-review/queue")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(queue.status_code, 200, queue.text)
        queued_verified = next(item for item in queue.json()["versions"] if item["version_id"] == version_id)
        self.assertTrue(queued_verified["answer_eligible"])
        self.assertTrue(queued_verified["source_permission_register_ready"])
        metrics = response.json()["metrics"]
        self.assertGreaterEqual(metrics["source_count"], 2)
        self.assertGreaterEqual(metrics["metadata_only_documents"], 1)
        self.assertEqual(metrics["verified_operational_versions"], verified_before + 1)
        self.assertGreaterEqual(metrics["permission_or_evidence_blocked_versions"], 1)
        self.assertEqual(metrics["central_permission_register_blocked_versions"], 0)
        self.assertGreaterEqual(metrics["central_permission_register_blocked_sources"], 1)

        # Coverage must use the latest event for every required review type,
        # matching the runtime answer-eligibility gate. Historical passes cannot
        # mask a later failure or an unattested latest lawyer approval.
        missing_review_events_before = metrics["versions_missing_review_event_evidence"]
        corpus.execute(
            "INSERT INTO legal_review_events(version_id,reviewer,review_type,outcome,notes,evidence_url) "
            "VALUES(?,?,?,?,?,?)",
            (version_id, "Synthetic reviewer", "source_authenticity", "fail",
             "Synthetic later withdrawal", "https://example.test/review-later-fail"),
        )
        corpus.execute(
            "INSERT INTO legal_review_events(version_id,reviewer,review_type,outcome,notes,evidence_url) "
            "VALUES(?,?,?,?,?,?)",
            (version_id, "Synthetic reviewer", "lawyer_approval", "pass",
             "Synthetic pass without counsel attestation", "https://example.test/review-unattested"),
        )
        corpus.commit()
        with patch.object(app, "_central_source_permission_evidence", side_effect=permission_record):
            latest_blocked = self.client.get("/legal-review/coverage")
        self.assertEqual(latest_blocked.status_code, 200, latest_blocked.text)
        self.assertEqual(
            latest_blocked.json()["metrics"]["versions_missing_review_event_evidence"],
            missing_review_events_before + 1,
        )

        # SQL-prefix lookalikes and credential-bearing HTTPS URLs are not
        # valid counsel evidence; the dashboard must use the strict shared gate.
        corpus.execute(
            "INSERT INTO legal_review_events(version_id,reviewer,review_type,outcome,notes,evidence_url) "
            "VALUES(?,?,?,?,?,?)",
            (version_id, "Synthetic counsel", "lawyer_approval", "pass",
             "[qualified_egyptian_counsel_attested:user_id=0] Synthetic malformed marker",
             "https://user:secret@example.test/review"),
        )
        corpus.commit()
        with patch.object(app, "_central_source_permission_evidence", side_effect=permission_record):
            malformed_latest = self.client.get("/legal-review/coverage")
        self.assertEqual(malformed_latest.status_code, 200, malformed_latest.text)
        self.assertEqual(
            malformed_latest.json()["metrics"]["versions_missing_review_event_evidence"],
            missing_review_events_before + 1,
        )

        # A later pass on both review types restores eligibility in the metric;
        # tie timestamps are resolved by the monotonically increasing event ID.
        corpus.execute(
            "INSERT INTO legal_review_events(version_id,reviewer,review_type,outcome,notes,evidence_url) "
            "VALUES(?,?,?,?,?,?)",
            (version_id, "Synthetic reviewer", "source_authenticity", "pass",
             "Synthetic current authenticity review", "https://example.test/review-latest-pass"),
        )
        corpus.execute(
            "INSERT INTO legal_review_events(version_id,reviewer,review_type,outcome,notes,evidence_url) "
            "VALUES(?,?,?,?,?,?)",
            (version_id, "Synthetic counsel", "lawyer_approval", "pass",
             "[qualified_egyptian_counsel_attested:user_id=42] Synthetic current counsel review",
             "https://example.test/review-current-counsel"),
        )
        corpus.commit()
        with patch.object(app, "_central_source_permission_evidence", side_effect=permission_record):
            latest_restored = self.client.get("/legal-review/coverage")
        self.assertEqual(latest_restored.status_code, 200, latest_restored.text)
        self.assertEqual(
            latest_restored.json()["metrics"]["versions_missing_review_event_evidence"],
            missing_review_events_before,
        )
        self.assertGreaterEqual(metrics["versions_not_eligible_for_verified_answers"], 1)
        self.assertTrue(any(
            row["practice_area"] == "labour"
            and row["documents_with_verified_operational_text"] == 1
            for row in response.json()["coverage_by_practice_area"]
        ))
        self.assertTrue(any(
            "not total Egyptian-law coverage" in warning
            for warning in response.json()["warnings"]
        ))

        blocked_before_unknown = metrics["permission_or_evidence_blocked_versions"]
        corpus.execute(
            "INSERT INTO legal_sources(source_id,authority,source_page,collection_permission,"
            "commercial_reuse_permission,permission_evidence) VALUES(?,?,?,?,?,?)",
            ("RUNTIME-COVERAGE-UNKNOWN-SOURCE", "Unknown permission authority",
             "https://example.test/unknown", "unknown", "unknown",
             "permission classification awaiting review"),
        )
        corpus.execute(
            "INSERT INTO legal_documents(document_id,source_id,official_title,document_type) "
            "VALUES(?,?,?,?)",
            ("RUNTIME-COVERAGE-UNKNOWN-DOC", "RUNTIME-COVERAGE-UNKNOWN-SOURCE",
             "Unknown permission instrument", "law"),
        )
        corpus.execute(
            "INSERT INTO legal_document_versions(document_id,version_label,source_url,legal_status,"
            "review_status,permission_evidence) VALUES(?,?,?,?,?,?)",
            ("RUNTIME-COVERAGE-UNKNOWN-DOC", "pending", "https://example.test/unknown-version",
             "under_review", "pending_source_review", "synthetic review evidence"),
        )
        corpus.commit()
        with patch.object(app, "_central_source_permission_evidence", side_effect=permission_record):
            after_unknown = self.client.get("/legal-review/coverage")
        self.assertEqual(after_unknown.status_code, 200, after_unknown.text)
        self.assertEqual(
            after_unknown.json()["metrics"]["permission_or_evidence_blocked_versions"],
            blocked_before_unknown + 1,
        )
        self.assertEqual(
            after_unknown.json()["metrics"]["central_permission_register_blocked_versions"],
            1,
        )


    def test_platform_reviewer_catalog_queries_are_metadata_only_and_bounded(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)

        control = self.runtime.control_connection()
        control.execute(
            "INSERT INTO user_platform_roles(user_id,role_name) VALUES(?,?)",
            (self.user_id, "platform_reviewer"),
        )
        control.commit()

        corpus = self.runtime.legal_corpus_connection()
        corpus.execute(
            "INSERT INTO legal_sources(source_id,authority,source_page,collection_permission,"
            "commercial_reuse_permission,permission_evidence) VALUES(?,?,?,?,?,?)",
            ("RUNTIME-CATALOG-SOURCE", "Runtime catalog authority",
             "https://example.test/catalog", "unknown", "unknown", "permission not yet verified"),
        )
        corpus.execute(
            "INSERT INTO legal_documents(document_id,source_id,official_title,document_type) "
            "VALUES(?,?,?,?)",
            ("RUNTIME-CATALOG-METADATA-DOC", "RUNTIME-CATALOG-SOURCE",
             "Metadata-only runtime instrument", "law"),
        )
        corpus.execute(
            "INSERT INTO legal_documents(document_id,source_id,official_title,document_type) "
            "VALUES(?,?,?,?)",
            ("RUNTIME-CATALOG-VERSIONED-DOC", "RUNTIME-CATALOG-SOURCE",
             "Versioned runtime instrument", "law"),
        )
        corpus.execute(
            "INSERT INTO legal_document_versions(document_id,version_label,source_url,legal_status,"
            "review_status,permission_evidence) VALUES(?,?,?,?,?,?)",
            ("RUNTIME-CATALOG-VERSIONED-DOC", "pending", "https://example.test/version",
             "under_review", "pending_source_review", "synthetic review evidence"),
        )
        corpus.commit()

        sources = self.client.get(
            "/legal-review/catalog?kind=sources&q=RUNTIME-CATALOG-SOURCE"
        )
        self.assertEqual(sources.status_code, 200, sources.text)
        self.assertTrue(any(
            row["source_id"] == "RUNTIME-CATALOG-SOURCE" for row in sources.json()["items"]
        ))
        self.assertTrue(all("text" not in row for row in sources.json()["items"]))

        metadata = self.client.get(
            "/legal-review/catalog?kind=metadata_only_documents&q=RUNTIME-CATALOG-METADATA-DOC"
        )
        self.assertEqual(metadata.status_code, 200, metadata.text)
        self.assertTrue(any(
            row["document_id"] == "RUNTIME-CATALOG-METADATA-DOC"
            and row["version_count"] == 0 for row in metadata.json()["items"]
        ))
        self.assertTrue(all("text" not in row for row in metadata.json()["items"]))

        permission_filtered = self.client.get(
            "/legal-review/catalog?kind=sources&collection_permission=unknown"
            "&commercial_reuse_permission=unknown"
        )
        self.assertEqual(permission_filtered.status_code, 200, permission_filtered.text)
        self.assertTrue(any(
            row["source_id"] == "RUNTIME-CATALOG-SOURCE"
            for row in permission_filtered.json()["items"]
        ))

        bounded = self.client.get("/legal-review/catalog?kind=sources&limit=999")
        self.assertEqual(bounded.status_code, 200, bounded.text)
        self.assertEqual(bounded.json()["limit"], 200)
        self.assertLessEqual(len(bounded.json()["items"]), 200)
        minimum = self.client.get("/legal-review/catalog?kind=sources&limit=0")
        self.assertEqual(minimum.status_code, 200, minimum.text)
        self.assertEqual(minimum.json()["limit"], 1)
        self.assertEqual(
            self.client.get("/legal-review/catalog?kind=unknown").status_code, 422
        )


    def test_date_models_reject_impossible_calendar_dates(self):
        cases = (
            (app.SessIn, {"date": "2024-02-29"}, "date", "2025-02-29"),
            (app.DeadlineIn, {"title": "deadline", "due": "2026-10-23"}, "due", "2026-02-30"),
            (app.TaskIn, {"title": "task", "due": "2024-02-29"}, "due", "2024-13-01"),
            (app.FeeIn, {"title": "fee", "amount": 10, "paid": 0, "due": "2024-02-29"}, "due", "2024-04-31"),
            (app.LawRow, {"law": "synthetic", "article": "1", "text": "synthetic", "as_of": "2024-02-29"}, "as_of", "2024-02-30"),
        )
        for model, valid_payload, field, impossible_date in cases:
            with self.subTest(model=model.__name__, impossible_date=impossible_date):
                model(**valid_payload)
                invalid_payload = dict(valid_payload)
                invalid_payload[field] = impossible_date
                with self.assertRaises(ValidationError):
                    model(**invalid_payload)


    def test_original_upload_preserves_duplicate_records_and_ocr_provenance(self):
        import hashlib

        payload = b"%PDF-1.4 synthetic scanned contract"
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "OCR provenance test"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        recognized = "Recognized scanned contract text requiring human review. " * 3

        with (
            patch.object(app, "extract", return_value=""),
            patch.object(app, "ocr_pdf_text", return_value=recognized),
            patch.object(app, "_mfa_required_for_user", return_value=False),
        ):
            first = self._mutate(
                "post", f"/cases/{case_id}/upload",
                files={"file": ("scan-one.pdf", payload, "application/pdf")},
            )
            second = self._mutate(
                "post", f"/cases/{case_id}/upload",
                files={"file": ("renamed-scan.pdf", payload, "application/pdf")},
            )

        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(second.status_code, 200, second.text)
        first_body, second_body = first.json(), second.json()
        self.assertTrue(first_body["ocr_used"])
        self.assertEqual(first_body["text_status"], "ocr_unverified")
        self.assertTrue(first_body["original_file_retained"])
        self.assertTrue(second_body["ocr_used"])
        self.assertEqual(second_body["text_status"], "ocr_unverified")
        self.assertTrue(second_body["original_file_retained"])
        self.assertNotEqual(first_body["document_id"], second_body["document_id"])
        expected_digest = hashlib.sha256(payload).hexdigest()
        self.assertEqual(first_body["sha256"], expected_digest)
        self.assertEqual(second_body["sha256"], expected_digest)

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        object_count_before_invalid_uploads = firm_db.execute(
            "SELECT COUNT(*) FROM original_file_objects"
        ).fetchone()[0]
        unsupported = self._mutate(
            "post", f"/cases/{case_id}/upload",
            files={"file": ("unsupported.txt", b"plain text", "text/plain")},
        )
        self.assertEqual(unsupported.status_code, 415, unsupported.text)
        malformed = self._mutate(
            "post", f"/cases/{case_id}/upload",
            files={"file": ("malformed.docx", b"not a docx",
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        self.assertEqual(malformed.status_code, 422, malformed.text)
        firm_db.commit()
        self.assertEqual(
            firm_db.execute("SELECT COUNT(*) FROM original_file_objects").fetchone()[0],
            object_count_before_invalid_uploads,
            "rejected uploads must not leave orphaned original-file metadata",
        )
        firm_db = self.runtime.firm_connection(context)
        document_ids = (first_body["document_id"], second_body["document_id"])
        documents = firm_db.execute(
            "SELECT id,firm_id,uploader_user_id,sha256,content_type,byte_size,extraction_status,ocr_used "
            "FROM docs WHERE id IN (?,?) ORDER BY id",
            document_ids,
        ).fetchall()
        self.assertEqual(len(documents), 2)
        for document in documents:
            self.assertEqual(document["firm_id"], self.firm_a)
            self.assertEqual(document["uploader_user_id"], self.user_id)
            self.assertEqual(document["sha256"], expected_digest)
            self.assertEqual(document["content_type"], "application/pdf")
            self.assertEqual(document["byte_size"], len(payload))
            self.assertEqual(document["extraction_status"], "ocr_unverified")
            self.assertEqual(document["ocr_used"], 1)

        objects = firm_db.execute(
            "SELECT document_id,object_id,sha256,byte_size,state FROM original_file_objects "
            "WHERE document_id IN (?,?) ORDER BY document_id",
            document_ids,
        ).fetchall()
        self.assertEqual(len(objects), 2)
        self.assertNotEqual(objects[0]["object_id"], objects[1]["object_id"])
        for stored in objects:
            self.assertEqual(stored["sha256"], expected_digest)
            self.assertEqual(stored["byte_size"], len(payload))
            self.assertEqual(stored["state"], "ready")
            self.assertTrue(
                (self.root / "objects" / f"firm-{self.firm_a}" /
                 (stored["object_id"] + ".blob")).is_file()
            )

    def test_document_review_rejects_stale_correction_digest(self):
        import hashlib

        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/cases", json={"title": "Stale review test"})
        self.assertEqual(created.status_code, 200, created.text)
        case_id = created.json()["id"]
        uploaded = self._mutate(
            "post", f"/cases/{case_id}/docs",
            json={"name": "review-source.txt", "text": "Original extracted text for stale review protection. " * 3},
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        document_id = uploaded.json()["document_id"]

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        firm_db = self.runtime.firm_connection(context)
        chunk = firm_db.execute(
            "SELECT id,text FROM chunks WHERE document_id=? AND case_id=? ORDER BY id LIMIT 1",
            (document_id, case_id),
        ).fetchone()
        self.assertIsNotNone(chunk)
        chunk_id = int(chunk["id"])
        original_text = str(chunk["text"])
        original_digest = hashlib.sha256(original_text.encode("utf-8")).hexdigest()
        # Release the read snapshot before the API worker commits the correction.
        firm_db.commit()
        review_url = f"/cases/{case_id}/docs/{document_id}/chunks/{chunk_id}/review"

        accepted = self._mutate(
            "post", review_url,
            json={
                "outcome": "corrected",
                "corrected_text": original_text + " [reviewed correction]",
                "expected_text_sha256": original_digest,
                "page_reference": "PDF p. 4",
                "note": "Current digest accepted",
            },
        )
        self.assertEqual(accepted.status_code, 201, accepted.text)
        self.assertTrue(accepted.json()["text_updated"])
        self.assertEqual(accepted.json()["page_reference"], "PDF p. 4")
        self.assertFalse(accepted.json()["source_page_verified"])

        review_audit = self.client.get(
            f"/cases/{case_id}/docs/{document_id}/reviews"
        )
        self.assertEqual(review_audit.status_code, 200, review_audit.text)
        self.assertEqual(review_audit.json()["reviews"][0]["outcome"], "corrected")
        self.assertEqual(review_audit.json()["reviews"][0]["original_text"], original_text)

        missing_digest = self._mutate(
            "post", review_url,
            json={"outcome": "corrected", "corrected_text": "correction without a digest"},
        )
        self.assertEqual(missing_digest.status_code, 422, missing_digest.text)

        stale = self._mutate(
            "post", review_url,
            json={
                "outcome": "corrected",
                "corrected_text": "stale overwrite must not win",
                "expected_text_sha256": original_digest,
            },
        )
        self.assertEqual(stale.status_code, 409, stale.text)
        # The test connection is thread-local and must refresh after API-worker commits.
        firm_db.commit()
        current = firm_db.execute(
            "SELECT text FROM chunks WHERE id=? AND document_id=? AND case_id=?",
            (chunk_id, document_id, case_id),
        ).fetchone()
        self.assertEqual(current["text"], original_text + " [reviewed correction]")
        reviews = firm_db.execute(
            "SELECT outcome,original_text,reviewed_text,page_reference FROM document_extraction_reviews "
            "WHERE chunk_id=? ORDER BY id",
            (chunk_id,),
        ).fetchall()
        self.assertEqual(len(reviews), 1)
        self.assertEqual(reviews[0]["outcome"], "corrected")
        self.assertEqual(reviews[0]["original_text"], original_text)
        self.assertEqual(reviews[0]["reviewed_text"], original_text + " [reviewed correction]")
        self.assertEqual(reviews[0]["page_reference"], "PDF p. 4")

    def test_research_followups_use_only_prior_user_questions_as_context(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        created = self._mutate("post", "/research/conversations", json={})
        self.assertEqual(created.status_code, 200, created.text)
        conversation_id = created.json()["id"]
        blank = self._mutate(
            "post", f"/research/conversations/{conversation_id}/messages",
            json={"message": "   "},
        )
        self.assertEqual(blank.status_code, 422, blank.text)
        captured = []
        assistant_only_sentinel = "ASSISTANT-ANSWER-MUST-NOT-BECOME-RETRIEVAL-CONTEXT"

        def fake_ask(db_conn, question, case_id=None, firm_id=None, as_of=None,
                     conversation_context=None, scope="verified_corpus"):
            captured.append({
                "question": question,
                "context": conversation_context,
                "firm_id": firm_id,
                "scope": scope,
            })
            return {
                "found": False,
                "answer": None,
                "sources": [],
                "message": assistant_only_sentinel,
                "scope": scope,
                "case_id": case_id,
            }

        with patch.object(app, "ask", side_effect=fake_ask):
            first = self._mutate(
                "post", f"/research/conversations/{conversation_id}/messages",
                json={"message": "مدة العقد"},
            )
            followup = self._mutate(
                "post", f"/research/conversations/{conversation_id}/messages",
                json={"message": "وهل توجد استثناءات؟"},
            )

        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(followup.status_code, 200, followup.text)
        self.assertEqual(len(captured), 2)
        self.assertEqual(captured[1]["question"], "وهل توجد استثناءات؟")
        self.assertIn("مدة العقد", captured[1]["context"] or "")
        self.assertNotIn(assistant_only_sentinel, captured[1]["context"] or "")
        self.assertEqual(captured[1]["firm_id"], self.firm_a)
        self.assertEqual(captured[1]["scope"], "verified_corpus")

        conversation = self.client.get(f"/research/conversations/{conversation_id}")
        self.assertEqual(conversation.status_code, 200, conversation.text)
        messages = conversation.json()["messages"]
        self.assertEqual([item["role"] for item in messages], ["user", "assistant", "user", "assistant"])
        self.assertEqual(messages[-1]["sources"], [])
        self.assertEqual(messages[-1]["content"], assistant_only_sentinel)
        self.assertTrue(any(
            item["id"] == conversation_id
            for item in self.client.get("/research/conversations").json()
        ))


    def test_case_lead_cannot_add_revoked_firm_member_to_case(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        switched = self._mutate(
            "post", "/session/firm", json={"firm_id": self.firm_a}
        )
        self.assertEqual(switched.status_code, 200, switched.text)
        case = self._mutate(
            "post", "/cases", json={"title": "Revoked member must stay excluded"}
        )
        self.assertEqual(case.status_code, 200, case.text)
        case_id = case.json()["id"]

        target_id = self.runtime.create_global_identity(
            "Revoked Case Member", "revoked-case-member@example.test",
            app.hash_pw(self.password),
        )
        self.runtime.create_firm_membership(target_id, self.firm_a, "lawyer")
        control = self.runtime.control_connection()
        control.execute(
            "UPDATE firm_memberships SET status='revoked',updated_at=datetime('now') "
            "WHERE user_id=? AND firm_id=?",
            (target_id, self.firm_a),
        )
        control.commit()

        denied = self._mutate(
            "post", f"/cases/{case_id}/members",
            json={"user_id": target_id, "role": "associate"},
        )
        self.assertEqual(denied.status_code, 404, denied.text)
        self.assertIn("active firm member not found", denied.text)

        context = self.runtime.resolve_context(self.user_id, self.firm_a)
        tenant = self.runtime.firm_connection(context)
        self.assertEqual(
            tenant.execute(
                "SELECT COUNT(*) FROM members WHERE case_id=? AND user_id=?",
                (case_id, target_id),
            ).fetchone()[0],
            0,
        )

    def test_revoked_firm_member_cannot_be_assigned_a_case_task(self):
        login = self.client.post(
            "/login", json={"email": "OWNER@example.test", "password": self.password}
        )
        self.assertEqual(login.status_code, 200, login.text)
        case = self._mutate(
            "post", "/cases", json={"title": "Revoked assignee must be excluded"}
        )
        self.assertEqual(case.status_code, 200, case.text)
        case_id = case.json()["id"]

        target_id = self.runtime.create_global_identity(
            "Revoked Task Assignee", "revoked-task-assignee@example.test",
            app.hash_pw(self.password),
        )
        self.runtime.create_firm_membership(target_id, self.firm_a, "lawyer")
        added = self._mutate(
            "post", f"/cases/{case_id}/members",
            json={"user_id": target_id, "role": "associate"},
        )
        self.assertEqual(added.status_code, 200, added.text)

        control = self.runtime.control_connection()
        control.execute(
            "UPDATE firm_memberships SET status='revoked',updated_at=datetime('now') "
            "WHERE user_id=? AND firm_id=?",
            (target_id, self.firm_a),
        )
        control.commit()

        denied = self._mutate(
            "post", f"/cases/{case_id}/tasks",
            json={"title": "Do not assign to revoked user", "assignee": target_id},
        )
        self.assertEqual(denied.status_code, 422, denied.text)
        self.assertIn("active firm member", denied.text)

        tenant = self.runtime.firm_connection(
            self.runtime.resolve_context(self.user_id, self.firm_a)
        )
        self.assertEqual(
            tenant.execute(
                "SELECT COUNT(*) FROM tasks WHERE case_id=? AND assignee=?",
                (case_id, target_id),
            ).fetchone()[0],
            0,
        )


    def test_research_context_does_not_cross_scope_or_matter_boundaries(self):
        prior = [
            {"role": "user", "content": "FIRM_ONLY_SENTINEL_8C21", "message_json": None},
            {
                "role": "assistant",
                "content": "Firm document answer",
                "message_json": json.dumps({"scope": "firm_documents", "case_id": 41}),
            },
            {"role": "user", "content": "VERIFIED_CORPUS_QUESTION", "message_json": None},
            {
                "role": "assistant",
                "content": "Verified corpus answer",
                "message_json": json.dumps({"scope": "verified_corpus", "case_id": None}),
            },
        ]
        corpus_context = app._research_context_for_scope(
            prior, "verified_corpus", None
        )
        self.assertEqual(corpus_context, "VERIFIED_CORPUS_QUESTION")
        self.assertNotIn("FIRM_ONLY_SENTINEL_8C21", corpus_context or "")

        same_matter = [
            {"role": "user", "content": "FIRST MATTER QUESTION", "message_json": None},
            {
                "role": "assistant",
                "content": "First answer",
                "message_json": json.dumps({"scope": "firm_documents", "case_id": 41}),
            },
            {"role": "user", "content": "FOLLOW-UP QUESTION", "message_json": None},
            {
                "role": "assistant",
                "content": "Follow-up answer",
                "message_json": json.dumps({"scope": "firm_documents", "case_id": 41}),
            },
        ]
        self.assertEqual(
            app._research_context_for_scope(same_matter, "firm_documents", 41),
            "FIRST MATTER QUESTION\nFOLLOW-UP QUESTION",
        )
        self.assertEqual(
            app._research_context_for_scope(same_matter, "firm_documents", 99),
            None,
        )

        legacy = [
            {"role": "user", "content": "LEGACY_PRIVATE_QUESTION", "message_json": None},
            {"role": "assistant", "content": "Legacy answer", "message_json": None},
        ]
        self.assertIsNone(
            app._research_context_for_scope(legacy, "verified_corpus", None)
        )

if __name__ == "__main__":
    unittest.main()
