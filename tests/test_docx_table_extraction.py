"""DOCX tables and headers/footers must be searchable, not silently omitted."""
import io
import unittest
import docx
import app


class DocxTableExtractionTests(unittest.TestCase):
    def test_preserves_paragraphs_tables_and_headers_footers_in_order(self):
        document = docx.Document()
        document.sections[0].header.paragraphs[0].text = "Court file 123/2026"
        document.add_paragraph("Parties and agreement terms")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Party"
        table.cell(0, 1).text = "Role"
        table.cell(1, 0).text = "أحمد"
        table.cell(1, 1).text = "Tenant"
        document.add_paragraph("Payment terms follow.")
        document.sections[0].footer.paragraphs[0].text = "Confidential"
        data = io.BytesIO()
        document.save(data)

        extracted = app.extract("agreement.docx", data.getvalue())
        self.assertLess(extracted.index("Parties and agreement terms"),
                        extracted.index("Party | Role"))
        self.assertIn("أحمد | Tenant", extracted)
        self.assertLess(extracted.index("أحمد | Tenant"),
                        extracted.index("Payment terms follow."))
        self.assertIn("[Header] Court file 123/2026", extracted)
        self.assertIn("[Footer] Confidential", extracted)


if __name__ == "__main__":
    unittest.main()
